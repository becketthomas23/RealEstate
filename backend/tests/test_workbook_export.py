from decimal import Decimal
from io import BytesIO
import unittest

from openpyxl import load_workbook

from backend.engines.operating import OperatingAssumptions
from backend.exports.workbook import build_underwriting_workbook
from backend.workflows.monte_carlo import MonteCarloAssumptions, SimulationVariable, calculate_monte_carlo
from backend.workflows.underwriting import UnderwritingAssumptions, calculate_underwriting


D = Decimal


class WorkbookExportTests(unittest.TestCase):
    def setUp(self) -> None:
        self.assumptions = UnderwritingAssumptions(
            purchase_price=D("8500000"), closing_cost_rate=D("0.02"), loan_to_value=D("0.60"),
            annual_interest_rate=D("0.06"), amortization_years=30, exit_cap_rate=D("0.065"),
            selling_cost_rate=D("0.02"), operating=OperatingAssumptions(
                current_potential_gross_income=D("1000000"), current_operating_expenses=D("400000"),
                vacancy_rate=D("0.05"), rent_growth_rate=D("0.03"), expense_growth_rate=D("0.025"), hold_period_years=5,
            ),
        )
        self.result = calculate_underwriting(self.assumptions)
        self.report = {
            "deal_name": "Maple Street Apartments",
            "version_label": "v2 · Base case",
            "analysis_basis": "solved maximum supportable price",
            "target_levered_irr": D("0.12"),
            "minimum_dscr": D("1.25"),
            "maximum_purchase_price": D("8500000"),
            "recommended_bid_floor": D("8075000"),
            "recommended_bid_ceiling": D("8500000"),
            "binding_constraint": "target_levered_irr",
        }
        self.detail = {
            "property": {"name": "Maple Street Apartments", "units": 48, "rentable_square_feet": D("42000"), "year_built": 1998},
            "rent_roll": [
                {"unit_type": "1BR", "units": 24, "average_square_feet": D("690"), "current_monthly_rent": D("1395"), "market_monthly_rent": D("1450")},
                {"unit_type": "2BR", "units": 24, "average_square_feet": D("980"), "current_monthly_rent": D("1795"), "market_monthly_rent": None},
            ],
            "other_income": [{"category": "Parking", "annual_amount": D("28800"), "note": None}],
            "operating_expenses": [
                {"category": "Taxes", "annual_amount": D("118000"), "note": "Reassessed"},
                {"category": "Insurance", "annual_amount": D("34000"), "note": None},
            ],
        }

    def test_workbook_contains_every_analyst_sheet_with_live_formulas(self) -> None:
        simulation = calculate_monte_carlo(
            MonteCarloAssumptions(
                underwriting=self.assumptions,
                variables=(SimulationVariable("rent_growth_rate", "medium", D("50")), SimulationVariable("exit_cap_rate", "medium", D("50"))),
                iterations=100,
                seed=7,
                target_levered_irr=D("0.12"),
                minimum_dscr=D("1.25"),
            )
        )

        output = build_underwriting_workbook(self.result, self.assumptions, report=self.report, detail=self.detail, simulation=simulation)

        book = load_workbook(BytesIO(output.getvalue()))
        self.assertEqual(
            book.sheetnames,
            ["Summary", "Pro Forma", "Monthly Cash Flow", "Debt", "Sale & Returns", "Sensitivity", "Risk Simulation", "Rent Roll & Expenses", "Risk Flags", "Assumptions"],
        )
        self.assertTrue(book.calculation.fullCalcOnLoad)

        summary = book["Summary"]
        self.assertEqual(summary["A1"].value, "Maple Street Apartments")
        labels = {row[0].value: row[1].value for row in summary.iter_rows(min_row=1, max_col=2) if row[0].value}
        self.assertEqual(labels["Analysis purchase price"], 8500000.0)
        self.assertEqual(labels["Maximum supportable price"], 8500000.0)
        self.assertAlmostEqual(labels["Price per unit"], 8500000 / 48)
        right = {row[0].value: row[1].value for row in summary.iter_rows(min_row=1, min_col=7, max_col=8) if row[0].value}
        self.assertAlmostEqual(right["Senior loan"], 5100000.0)
        self.assertAlmostEqual(right["Equity"], 8500000 * 1.02 - 5100000)

        pro_forma = book["Pro Forma"]
        formulas = [cell.value for row in pro_forma.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")]
        self.assertTrue(any(value.startswith("=IRR(") for value in formulas))
        self.assertTrue(any("+" in value for value in formulas))
        line_items = [row[0].value for row in pro_forma.iter_rows(min_col=1, max_col=1) if row[0].value]
        for expected in ("Potential gross income", "Net operating income", "Total debt service", "Levered cash flow", "Unlevered cash flow (NOI + net sale before debt)", "Levered IRR"):
            self.assertIn(expected, line_items)

        monthly = book["Monthly Cash Flow"]
        self.assertEqual(monthly.max_row, 5 + 60)  # title, subtitle, blank, header, 60 months, total row
        debt = book["Debt"]
        debt_labels = {row[0].value for row in debt.iter_rows(min_col=1, max_col=1) if row[0].value}
        self.assertIn("Loan constant (annual debt service / loan)", debt_labels)
        self.assertIn("Monthly amortization schedule", debt_labels)

        sensitivity = book["Sensitivity"]
        titles = [row[0].value for row in sensitivity.iter_rows(min_col=1, max_col=1) if isinstance(row[0].value, str) and "×" in row[0].value]
        self.assertGreaterEqual(len(titles), 5)
        self.assertTrue(any(title.startswith("Maximum supportable price") for title in titles))

        risk = book["Risk Simulation"]
        risk_labels = {row[0].value for row in risk.iter_rows(min_col=1, max_col=1) if row[0].value}
        self.assertIn("Probability of a loss (IRR below 0%)", risk_labels)
        self.assertIn("Driver impact (tornado data)", risk_labels)

        detail = book["Rent Roll & Expenses"]
        detail_labels = {row[0].value for row in detail.iter_rows(min_col=1, max_col=1) if row[0].value}
        self.assertIn("Total / weighted", detail_labels)
        self.assertIn("Total operating expenses", detail_labels)

    def test_workbook_without_optional_inputs_skips_optional_sheets(self) -> None:
        all_cash = UnderwritingAssumptions(
            purchase_price=D("8500000"), closing_cost_rate=D("0.02"), loan_to_value=D("0"),
            annual_interest_rate=D("0.06"), amortization_years=30, exit_cap_rate=D("0.065"),
            selling_cost_rate=D("0.02"), operating=self.assumptions.operating,
        )

        output = build_underwriting_workbook(calculate_underwriting(all_cash), all_cash, include_sensitivity=False)

        book = load_workbook(BytesIO(output.getvalue()))
        self.assertEqual(book.sheetnames, ["Summary", "Pro Forma", "Monthly Cash Flow", "Debt", "Sale & Returns", "Risk Flags", "Assumptions"])
        self.assertEqual(book["Debt"]["A4"].value, "All-cash acquisition: no debt is modeled.")
