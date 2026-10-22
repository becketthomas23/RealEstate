from datetime import date
from decimal import Decimal
from io import BytesIO
import unittest

from openpyxl import Workbook, load_workbook

from backend.workbook_import import parse_underwriting_workbook, underwriting_workbook_template


D = Decimal


def _save(book: Workbook) -> bytes:
    output = BytesIO()
    book.save(output)
    return output.getvalue()


class WorkbookTemplateTests(unittest.TestCase):
    def test_template_round_trips_through_the_parser_with_rolled_up_income_and_expenses(self) -> None:
        template = underwriting_workbook_template().getvalue()

        book = load_workbook(BytesIO(template))
        self.assertEqual(book.sheetnames, ["Instructions", "Assumptions", "Rent Roll", "Other Income", "Operating Expenses"])

        imported = parse_underwriting_workbook(template)

        underwriting = imported.underwriting
        self.assertEqual(underwriting.purchase_price, D("8500000"))
        self.assertEqual(underwriting.closing_cost_rate, D("0.02"))
        self.assertEqual(underwriting.loan_to_value, D("0.6"))
        self.assertEqual(underwriting.operating.hold_period_years, 5)
        # Rent roll: (8x1150 + 24x1395 + 16x1795) x 12 = 856,800, plus 52,800 of other income.
        self.assertEqual(underwriting.operating.current_potential_gross_income, D("909600"))
        self.assertEqual(imported.sources["current_potential_gross_income"], "rolled up from Rent Roll and Other Income")
        # Operating expenses sheet sums to 405,400.
        self.assertEqual(underwriting.operating.current_operating_expenses, D("405400"))
        self.assertFalse(underwriting.operating.apply_growth_in_year_one)
        detail = imported.supporting_detail
        self.assertEqual(detail.property.name, "Maple Street Apartments")
        self.assertEqual(detail.property.units, 48)
        self.assertEqual(len(detail.rent_roll), 3)
        self.assertEqual(len(detail.operating_expenses), 10)
        self.assertEqual(imported.version["label"], "Broker OM base case")
        self.assertEqual(imported.version["effective_date"], "2026-09-30")
        self.assertEqual(imported.version["confidence_level"], "medium")
        self.assertEqual(imported.version["property_type"], "Multifamily")
        self.assertEqual(imported.targets["target_levered_irr"], D("0.12"))
        self.assertEqual(imported.targets["minimum_dscr"], D("1.25"))
        self.assertTrue(any("loss-to-lease" in warning for warning in imported.warnings))


class WorkbookParserTests(unittest.TestCase):
    def _minimal_book(self, **overrides) -> Workbook:
        book = Workbook()
        sheet = book.active
        sheet.title = "Inputs"
        rows = {
            "Asking Price": 1000000,
            "Closing Costs": "2%",
            "Hold Period": 5,
            "Terminal Cap Rate": 10,
            "Cost of Sale": 0.02,
            "GPR": 200000,
            "OpEx": 50000,
            "Vacancy": 10,
            "Rent Growth": 0,
            "Expense Growth": 0,
            "LTV": 60,
            "Note Rate": 0,
            "Amort": 5,
            "As of": date(2026, 10, 1),
            "Scenario": "Downside",
        }
        rows.update(overrides)
        sheet.append(["Label", "Value"])
        for label, value in rows.items():
            if value is not None:
                sheet.append([label, value])
        return book

    def test_matches_aliases_and_accepts_percent_as_whole_numbers_strings_or_decimals(self) -> None:
        imported = parse_underwriting_workbook(_save(self._minimal_book()))

        underwriting = imported.underwriting
        self.assertEqual(underwriting.purchase_price, D("1000000"))
        self.assertEqual(underwriting.closing_cost_rate, D("0.02"))
        self.assertEqual(underwriting.exit_cap_rate, D("0.1"))
        self.assertEqual(underwriting.selling_cost_rate, D("0.02"))
        self.assertEqual(underwriting.loan_to_value, D("0.6"))
        self.assertEqual(underwriting.annual_interest_rate, D("0"))
        self.assertEqual(underwriting.amortization_years, 5)
        self.assertEqual(underwriting.operating.vacancy_rate, D("0.1"))
        self.assertEqual(imported.version["effective_date"], "2026-10-01")
        self.assertEqual(imported.version["label"], "Downside")
        self.assertEqual(imported.supporting_detail.rent_roll, [])
        self.assertIsNone(imported.supporting_detail.property)

    def test_reports_missing_required_inputs_by_label(self) -> None:
        with self.assertRaisesRegex(ValueError, "Exit cap rate"):
            parse_underwriting_workbook(_save(self._minimal_book(**{"Terminal Cap Rate": None})))

    def test_requires_income_or_a_rent_roll(self) -> None:
        with self.assertRaisesRegex(ValueError, "Rent Roll"):
            parse_underwriting_workbook(_save(self._minimal_book(GPR=None)))

    def test_warns_when_entered_income_disagrees_with_the_rent_roll(self) -> None:
        book = self._minimal_book()
        roll = book.create_sheet("Rent Roll")
        roll.append(["Unit type", "Units", "Average square feet", "Current monthly rent", "Market monthly rent"])
        roll.append(["1BR", 10, 700, 1000, 1100])

        imported = parse_underwriting_workbook(_save(book))

        self.assertEqual(imported.underwriting.operating.current_potential_gross_income, D("200000"))
        self.assertTrue(any("differs from the Rent Roll" in warning for warning in imported.warnings))
        self.assertEqual(imported.supporting_detail.property.units, 10)
        self.assertEqual(imported.supporting_detail.property.rentable_square_feet, D("7000"))

    def test_rejects_a_file_that_is_not_a_workbook(self) -> None:
        with self.assertRaisesRegex(ValueError, "readable"):
            parse_underwriting_workbook(b"not a workbook")
