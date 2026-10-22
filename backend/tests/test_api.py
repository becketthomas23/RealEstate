import unittest
from io import BytesIO
from pathlib import Path
from tempfile import TemporaryDirectory

from fastapi.testclient import TestClient
from openpyxl import load_workbook

import backend.api.main as api_main
from backend.api.main import app
from backend.persistence import Database


client = TestClient(app)

UNDERWRITING_REQUEST = {
    "purchase_price": "1000000",
    "closing_cost_rate": "0.02",
    "loan_to_value": "0.60",
    "annual_interest_rate": "0",
    "amortization_years": 5,
    "exit_cap_rate": "0.10",
    "selling_cost_rate": "0.02",
    "operating": {
        "current_potential_gross_income": "200000",
        "current_operating_expenses": "50000",
        "vacancy_rate": "0.10",
        "rent_growth_rate": "0",
        "expense_growth_rate": "0",
        "hold_period_years": 5,
    },
}


class ApiTests(unittest.TestCase):
    def test_health_check(self) -> None:
        response = client.get("/health")

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json(), {"status": "ok"})

    def test_local_frontend_origin_is_allowed(self) -> None:
        response = client.options(
            "/v1/underwriting",
            headers={
                "Origin": "http://localhost:3000",
                "Access-Control-Request-Method": "POST",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["access-control-allow-origin"], "http://localhost:3000")

    def test_underwriting_endpoint_returns_precise_financial_values(self) -> None:
        response = client.post("/v1/underwriting", json=UNDERWRITING_REQUEST)

        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["terminal_noi"], "130000.00")
        self.assertEqual(body["going_in_cap_rate"], "0.13")
        self.assertEqual(body["debt_yield"], "0.2166666666666666666666666667")
        self.assertEqual(body["annual_operating_cash_flows"][-1], "10000.00")
        self.assertEqual(body["annual_cash_flows"][-1], "1284000.00")
        self.assertEqual(body["returns"]["cash_on_cash_by_year"][-1], "0.02380952380952380952380952381")
        self.assertEqual(body["returns"]["minimum_dscr"], "1.083333333333333333333333333")
        self.assertEqual([flag["code"] for flag in body["risk_flags"]], ["thin_dscr", "exit_cap_compression"])

    def test_underwriting_accepts_a_rent_growth_schedule_and_returns_monthly_cash_flow(self) -> None:
        request = {
            **UNDERWRITING_REQUEST,
            "operating": {
                **UNDERWRITING_REQUEST["operating"],
                "annual_rent_growth_rates": ["0.10", "0"],
            },
        }

        response = client.post("/v1/underwriting", json=request)

        self.assertEqual(response.status_code, 200)
        body = response.json()
        # Year 1 is the entered current year; the 10% step applies in Year 2 and 0% after.
        self.assertEqual(body["operating_projection"][0]["potential_gross_income"], "200000")
        self.assertEqual(body["operating_projection"][1]["potential_gross_income"], "220000.00")
        self.assertEqual(body["operating_projection"][2]["potential_gross_income"], "220000.00")
        self.assertEqual(len(body["monthly_operating_projection"]), 60)
        self.assertEqual(body["monthly_operating_projection"][0]["net_operating_income"], "10833.33333333333333333333333")
        self.assertEqual(body["monthly_operating_projection"][12]["net_operating_income"], "12333.33333333333333333333333")

    def test_underwriting_supports_an_all_cash_deal_and_exports_it_without_debt_metrics(self) -> None:
        request = {**UNDERWRITING_REQUEST, "loan_to_value": "0"}

        response = client.post("/v1/underwriting", json=request)
        export = client.post("/v1/exports/memo", json=request)

        self.assertEqual(response.status_code, 200)
        self.assertIsNone(response.json()["debt_yield"])
        self.assertIsNone(response.json()["returns"]["minimum_dscr"])
        self.assertEqual(response.json()["disposition"]["loan_payoff"], "0")
        self.assertEqual(export.status_code, 200)
        self.assertEqual(export.headers["content-type"], "application/pdf")

    def test_maximum_price_endpoint_returns_dscr_limited_price(self) -> None:
        response = client.post(
            "/v1/maximum-price",
            json={
                "underwriting": UNDERWRITING_REQUEST,
                "target_levered_irr": "0.10",
                "minimum_dscr": "1.00",
                "minimum_purchase_price": "500000",
                "maximum_purchase_price": "1500000",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertLess(abs(float(response.json()["maximum_purchase_price"]) - 1083333.33), 0.01)
        self.assertEqual(response.json()["recommended_bid_ceiling"], response.json()["maximum_purchase_price"])

    def test_sensitivity_endpoint_returns_one_way_price_cases(self) -> None:
        response = client.post(
            "/v1/sensitivities",
            json={
                "underwriting": UNDERWRITING_REQUEST,
                "target_levered_irr": "0.10",
                "minimum_dscr": "1.00",
                "minimum_purchase_price": "500000",
                "maximum_purchase_price": "1500000",
            },
        )

        self.assertEqual(response.status_code, 200)
        dimensions = response.json()["dimensions"]
        self.assertEqual([dimension["key"] for dimension in dimensions], ["rent_growth_rate", "exit_cap_rate", "annual_interest_rate"])
        self.assertEqual(dimensions[0]["points"][1]["label"], "Base")
        self.assertGreater(
            float(dimensions[0]["points"][2]["maximum_purchase_price"]),
            float(dimensions[0]["points"][0]["maximum_purchase_price"]),
        )

    def test_advanced_sensitivity_endpoint_mixes_three_selected_inputs(self) -> None:
        response = client.post(
            "/v1/advanced-sensitivities",
            json={
                "underwriting": UNDERWRITING_REQUEST,
                "dimensions": [
                    {"key": "rent_growth_rate", "values": ["0", "0.03"]},
                    {"key": "exit_cap_rate", "values": ["0.08", "0.10"]},
                    {"key": "vacancy_rate", "values": ["0.05", "0.10"]},
                ],
                "metric": "levered_irr",
            },
        )

        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["metric"], "levered_irr")
        self.assertEqual(len(body["scenarios"]), 8)
        self.assertEqual(body["scenarios"][0]["values"], {"rent_growth_rate": "0", "exit_cap_rate": "0.08", "vacancy_rate": "0.05"})

    def test_monte_carlo_endpoint_returns_confidence_aware_summary(self) -> None:
        response = client.post(
            "/v1/monte-carlo",
            json={
                "underwriting": UNDERWRITING_REQUEST,
                "variables": [
                    {"key": "rent_growth_rate", "confidence_level": "medium", "confidence_percentage": "50"},
                    {"key": "exit_cap_rate", "confidence_level": "high", "confidence_percentage": "80"},
                    {"key": "vacancy_rate", "confidence_level": "low", "confidence_percentage": "20"},
                ],
                "iterations": 100,
                "seed": 7,
                "target_levered_irr": "0.10",
            },
        )

        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["iterations"], 100)
        self.assertEqual(body["levered_irr"]["metric"], "levered_irr")
        self.assertIsNotNone(body["levered_irr"]["probability_meeting_target"])

    def test_market_calibration_preview_accepts_aggregate_history_without_persisting_it(self) -> None:
        rows = "\n".join(
            f"202{year}-Q1,Multifamily,$10M-$25M,{0.02 + year / 1000:.3f},{0.060 - year / 1000:.3f},{0.07 - year / 2000:.3f},{0.02 + year / 1500:.3f},{0.04 + year / 1200:.3f}"
            for year in range(8)
        )
        contents = (
            "period,property_type,value_bucket,rent_growth_rate,exit_cap_rate,vacancy_rate,expense_growth_rate,annual_interest_rate\n"
            + rows
        ).encode()

        response = client.post(
            "/v1/market-calibrations/preview",
            data={
                "property_type": "Multifamily",
                "value_bucket": "$10M-$25M",
                "source": "Synthetic test market history",
                "source_as_of_date": "2026-10-21",
            },
            files={"file": ("market-history.csv", contents, "text/csv")},
        )

        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["context"]["observations"], 8)
        self.assertEqual(body["context"]["source"], "Synthetic test market history")
        self.assertEqual(body["context"]["source_as_of_date"], "2026-10-21")
        self.assertEqual(len(body["correlations"]), 10)
        self.assertEqual(len(body["driver_ranges"]), 5)
        self.assertTrue(body["warnings"])

        template = client.get("/v1/market-calibrations/template")
        self.assertEqual(template.status_code, 200)
        self.assertIn("property_type", template.text)

    def test_waterfall_endpoint_uses_the_underwriting_cash_flows(self) -> None:
        response = client.post(
            "/v1/waterfall",
            json={
                "underwriting": UNDERWRITING_REQUEST,
                "preferred_return_rate": "0.08",
                "sponsor_residual_share": "0.20",
            },
        )

        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(len(body["years"]), 5)
        self.assertEqual(body["years"][0]["property_cash_flow"], "10000.00")
        self.assertGreater(float(body["lp_irr"]), 0)

    def test_construction_loan_endpoint_returns_monthly_draw_schedule(self) -> None:
        response = client.post(
            "/v1/construction-loan",
            json={
                "monthly_project_costs": ["100", "200"],
                "loan_to_cost": "0.70",
                "annual_interest_rate": "0.12",
                "origination_fee_rate": "0.01",
            },
        )

        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.json()["ending_loan_balance"], "210.7000")

    def test_development_endpoint_connects_construction_refinance_and_stabilized_returns(self) -> None:
        response = client.post(
            "/v1/development",
            json={
                "construction": {
                    "monthly_project_costs": ["100", "200"],
                    "loan_to_cost": "0.70",
                    "annual_interest_rate": "0.12",
                    "origination_fee_rate": "0.01",
                },
                "operating": {
                    "current_potential_gross_income": "2000",
                    "current_operating_expenses": "500",
                    "vacancy_rate": "0.10",
                    "rent_growth_rate": "0",
                    "expense_growth_rate": "0",
                    "hold_period_years": 5,
                },
                "stabilization_cap_rate": "0.10",
                "permanent_loan_to_value": "0.60",
                "permanent_interest_rate": "0",
                "permanent_amortization_years": 30,
                "exit_cap_rate": "0.10",
                "selling_cost_rate": "0.02",
                "refinance_cost_rate": "0.01",
            },
        )

        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertEqual(body["construction"]["ending_loan_balance"], "210.7000")
        self.assertEqual(body["refinance"]["net_refinance_proceeds"], "7511.3000")
        self.assertEqual(body["annual_cash_flows"][0], "8551.300000000000000000000000")

    def test_underwriting_csv_template_and_import_validate_a_complete_assumption_row(self) -> None:
        template = client.get("/v1/imports/underwriting-template")

        self.assertEqual(template.status_code, 200)
        self.assertEqual(template.headers["content-type"], "text/csv; charset=utf-8")
        imported = client.post(
            "/v1/imports/underwriting",
            files={"file": ("underwriting.csv", template.content, "text/csv")},
        )

        self.assertEqual(imported.status_code, 200)
        self.assertEqual(imported.json()["underwriting"]["purchase_price"], "1000000")
        self.assertEqual(imported.json()["underwriting"]["operating"]["hold_period_years"], 5)

    def test_underwriting_csv_import_rejects_an_unrecognized_header(self) -> None:
        response = client.post(
            "/v1/imports/underwriting",
            files={"file": ("underwriting.csv", b"price\n1000000\n", "text/csv")},
        )

        self.assertEqual(response.status_code, 422)
        self.assertIn("headers", response.json()["detail"])

    def test_export_endpoints_return_real_excel_and_pdf_files(self) -> None:
        workbook = client.post("/v1/exports/excel", json=UNDERWRITING_REQUEST)
        memo = client.post("/v1/exports/memo", json=UNDERWRITING_REQUEST)

        self.assertEqual(workbook.status_code, 200)
        self.assertEqual(workbook.headers["content-type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertTrue(workbook.content.startswith(b"PK"))
        exported_book = load_workbook(BytesIO(workbook.content), data_only=True)
        risk_sheet = exported_book["Risk Flags"]
        self.assertEqual(risk_sheet[4][0].value, "Code")
        self.assertEqual(risk_sheet[5][0].value, "thin_dscr")
        assumption_sheet = exported_book["Assumptions"]
        assumption_labels = {row[0].value: row[1].value for row in assumption_sheet.iter_rows(min_col=1, max_col=2) if row[0].value}
        self.assertEqual(assumption_labels["Analysis purchase price"], 1000000.0)
        self.assertEqual(assumption_labels["Hold period"], 5)
        self.assertIn("Pro Forma", exported_book.sheetnames)
        self.assertIn("Sensitivity", exported_book.sheetnames)
        self.assertEqual(memo.status_code, 200)
        self.assertEqual(memo.headers["content-type"], "application/pdf")
        self.assertTrue(memo.content.startswith(b"%PDF"))

    def test_api_returns_validation_error_for_unbracketed_solver_range(self) -> None:
        response = client.post(
            "/v1/maximum-price",
            json={
                "underwriting": UNDERWRITING_REQUEST,
                "target_levered_irr": "0.10",
                "minimum_dscr": "0.50",
                "minimum_purchase_price": "500000",
                "maximum_purchase_price": "600000",
            },
        )

        self.assertEqual(response.status_code, 422)
        self.assertIn("higher bound", response.json()["detail"])


class PersistedUnderwritingApiTests(unittest.TestCase):
    """Exercise durable deal history separately from the stateless calculator endpoints."""

    def setUp(self) -> None:
        self.temp_directory = TemporaryDirectory()
        self.original_database = api_main.database
        database_path = Path(self.temp_directory.name) / "underwriting.db"
        api_main.database = Database(f"sqlite:///{database_path.as_posix()}")
        self.client_context = TestClient(api_main.app)
        self.client = self.client_context.__enter__()

    def tearDown(self) -> None:
        self.client_context.__exit__(None, None, None)
        api_main.database.dispose()
        api_main.database = self.original_database
        self.temp_directory.cleanup()

    def test_versions_inputs_and_results_are_persisted(self) -> None:
        deal = self.client.post("/v1/deals", json={"name": "Maple Apartments"})

        self.assertEqual(deal.status_code, 201)
        deal_id = deal.json()["id"]
        assumption_set = self.client.post(
            f"/v1/deals/{deal_id}/assumption-sets",
            json={
                "label": "IC base case",
                "source": "Broker OM",
                "effective_date": "2026-09-24",
                "confidence_level": "medium",
                "property_type": "Multifamily",
                "value_bucket": "$5M-$10M",
                "underwriting": UNDERWRITING_REQUEST,
            },
        )

        self.assertEqual(assumption_set.status_code, 201)
        assumptions = assumption_set.json()
        self.assertEqual(assumptions["version"], 1)
        self.assertEqual(assumptions["underwriting"]["purchase_price"], "1000000")
        self.assertEqual(assumptions["property_type"], "Multifamily")
        self.assertEqual(assumptions["value_bucket"], "$5M-$10M")

        calculation = self.client.post(f"/v1/assumption-sets/{assumptions['id']}/underwriting")

        self.assertEqual(calculation.status_code, 201)
        body = calculation.json()
        self.assertEqual(body["calculation_run"]["status"], "completed")
        self.assertEqual(body["result"]["annual_cash_flows"][-1], "1284000.00")

        saved_run = self.client.get(f"/v1/calculation-runs/{body['calculation_run']['id']}")
        self.assertEqual(saved_run.status_code, 200)
        self.assertEqual(saved_run.json()["result"], body["result"])

        sensitivity = self.client.post(
            f"/v1/assumption-sets/{assumptions['id']}/sensitivities",
            json={
                "target_levered_irr": "0.10",
                "minimum_dscr": "1.00",
                "minimum_purchase_price": "500000",
                "maximum_purchase_price": "1500000",
            },
        )
        self.assertEqual(sensitivity.status_code, 201)
        self.assertEqual(sensitivity.json()["calculation_run"]["calculation_type"], "sensitivity")
        self.assertEqual(len(sensitivity.json()["result"]["dimensions"]), 3)

        run_history = self.client.get(f"/v1/assumption-sets/{assumptions['id']}/calculation-runs")
        self.assertEqual(run_history.status_code, 200)
        self.assertEqual(
            [run["id"] for run in run_history.json()],
            [sensitivity.json()["calculation_run"]["id"], body["calculation_run"]["id"]],
        )

        next_version = self.client.post(
            f"/v1/deals/{deal_id}/assumption-sets",
            json={"underwriting": {**UNDERWRITING_REQUEST, "purchase_price": "1100000"}},
        )
        self.assertEqual(next_version.status_code, 201)
        self.assertEqual(next_version.json()["version"], 2)
        self.assertEqual(next_version.json()["property_type"], "Multifamily")
        self.assertEqual(next_version.json()["value_bucket"], "Custom")

        detail = self.client.get(f"/v1/deals/{deal_id}")
        self.assertEqual(detail.status_code, 200)
        self.assertEqual([item["version"] for item in detail.json()["assumption_sets"]], [2, 1])

    def test_risk_and_waterfall_runs_are_saved_against_the_input_version(self) -> None:
        deal = self.client.post("/v1/deals", json={"name": "Auditable advanced analysis"})
        self.assertEqual(deal.status_code, 201)
        assumptions = self.client.post(
            f"/v1/deals/{deal.json()['id']}/assumption-sets",
            json={"underwriting": UNDERWRITING_REQUEST},
        )
        self.assertEqual(assumptions.status_code, 201)
        assumption_set_id = assumptions.json()["id"]

        simulation = self.client.post(
            f"/v1/assumption-sets/{assumption_set_id}/monte-carlo",
            json={
                "variables": [{"key": "rent_growth_rate", "confidence_level": "medium", "confidence_percentage": "50"}],
                "iterations": 100,
                "seed": 7,
                "target_levered_irr": "0.10",
            },
        )
        waterfall = self.client.post(
            f"/v1/assumption-sets/{assumption_set_id}/waterfall",
            json={"preferred_return_rate": "0.08", "sponsor_residual_share": "0.20"},
        )

        self.assertEqual(simulation.status_code, 201)
        self.assertEqual(simulation.json()["calculation_run"]["calculation_type"], "monte_carlo")
        self.assertEqual(waterfall.status_code, 201)
        self.assertEqual(waterfall.json()["calculation_run"]["calculation_type"], "waterfall")
        history = self.client.get(f"/v1/assumption-sets/{assumption_set_id}/calculation-runs")
        self.assertEqual(history.status_code, 200)
        self.assertEqual([run["calculation_type"] for run in history.json()], ["waterfall", "monte_carlo"])

    def test_scenario_comparison_evaluates_each_immutable_deal_version(self) -> None:
        deal = self.client.post("/v1/deals", json={"name": "Scenario comparison"})
        self.assertEqual(deal.status_code, 201)
        deal_id = deal.json()["id"]
        base = self.client.post(
            f"/v1/deals/{deal_id}/assumption-sets",
            json={"label": "Base", "underwriting": UNDERWRITING_REQUEST},
        )
        downside = self.client.post(
            f"/v1/deals/{deal_id}/assumption-sets",
            json={
                "label": "Downside",
                "underwriting": {
                    **UNDERWRITING_REQUEST,
                    "operating": {**UNDERWRITING_REQUEST["operating"], "vacancy_rate": "0.15"},
                },
            },
        )
        self.assertEqual(base.status_code, 201)
        self.assertEqual(downside.status_code, 201)

        comparison = self.client.post(
            f"/v1/deals/{deal_id}/scenario-comparison",
            json={
                "target_levered_irr": "0.10",
                "minimum_dscr": "1.00",
                "minimum_purchase_price": "500000",
                "maximum_purchase_price": "1500000",
            },
        )

        self.assertEqual(comparison.status_code, 200)
        scenarios = comparison.json()["scenarios"]
        self.assertEqual([scenario["label"] for scenario in scenarios], ["Downside", "Base"])
        self.assertIsNone(scenarios[0]["error"])
        self.assertLess(float(scenarios[0]["maximum_purchase_price"]), float(scenarios[1]["maximum_purchase_price"]))
        self.assertEqual(scenarios[0]["recommended_bid_ceiling"], scenarios[0]["maximum_purchase_price"])


class ConventionAndAnalysisPriceApiTests(unittest.TestCase):
    """Cover the behaviours that make every panel describe the same deal."""

    def setUp(self) -> None:
        self.temp_directory = TemporaryDirectory()
        self.original_database = api_main.database
        database_path = Path(self.temp_directory.name) / "underwriting.db"
        api_main.database = Database(f"sqlite:///{database_path.as_posix()}")
        self.client_context = TestClient(api_main.app)
        self.client = self.client_context.__enter__()

    def tearDown(self) -> None:
        self.client_context.__exit__(None, None, None)
        api_main.database.dispose()
        api_main.database = self.original_database
        self.temp_directory.cleanup()

    def test_maximum_price_brackets_itself_and_names_the_binding_constraint(self) -> None:
        response = client.post(
            "/v1/maximum-price",
            json={"underwriting": UNDERWRITING_REQUEST, "target_levered_irr": "0.10", "minimum_dscr": "1.00"},
        )

        self.assertEqual(response.status_code, 200)
        body = response.json()
        self.assertLess(abs(float(body["maximum_purchase_price"]) - 1083333.33), 0.02)
        self.assertEqual(body["binding_constraint"], "minimum_dscr")
        self.assertEqual(body["underwriting"]["purchase_price"], body["maximum_purchase_price"])
        self.assertIn("initial_equity", body["underwriting"])

    def test_saved_risk_and_waterfall_runs_can_use_the_solved_analysis_price(self) -> None:
        deal = self.client.post("/v1/deals", json={"name": "Analysis price"})
        assumptions = self.client.post(f"/v1/deals/{deal.json()['id']}/assumption-sets", json={"underwriting": UNDERWRITING_REQUEST})
        assumption_set_id = assumptions.json()["id"]
        solved = self.client.post(
            f"/v1/assumption-sets/{assumption_set_id}/maximum-price",
            json={"target_levered_irr": "0.10", "minimum_dscr": "1.00"},
        )
        self.assertEqual(solved.status_code, 201)
        analysis_price = solved.json()["result"]["maximum_purchase_price"]

        simulation = self.client.post(
            f"/v1/assumption-sets/{assumption_set_id}/monte-carlo",
            json={
                "variables": [{"key": "vacancy_rate", "confidence_level": "high", "confidence_percentage": "100"}],
                "iterations": 100,
                "seed": 1,
                "target_levered_irr": "0.10",
                "purchase_price": analysis_price,
            },
        )
        waterfall = self.client.post(
            f"/v1/assumption-sets/{assumption_set_id}/waterfall",
            json={"preferred_return_rate": "0.08", "sponsor_residual_share": "0.20", "purchase_price": analysis_price},
        )

        self.assertEqual(simulation.status_code, 201)
        self.assertEqual(simulation.json()["calculation_run"]["request"]["purchase_price"], analysis_price)
        # The simulation describes the solved price: its median IRR sits near the
        # IRR at that price, well below the 26.5% the 1,000,000 reference price gives.
        solved_irr = float(solved.json()["result"]["underwriting"]["returns"]["levered_irr"])
        median_irr = float(simulation.json()["result"]["levered_irr"]["percentile_50"])
        self.assertLess(median_irr, 0.2646)
        self.assertLess(abs(median_irr - solved_irr), 0.02)
        self.assertEqual(waterfall.status_code, 201)
        self.assertEqual(waterfall.json()["calculation_run"]["request"]["purchase_price"], analysis_price)
        self.assertTrue(waterfall.json()["result"]["pay_preferred_return_before_capital"])
        # The waterfall must distribute exactly the deal's cash flow.
        lp = float(waterfall.json()["result"]["total_lp_distributed"])
        sponsor = float(waterfall.json()["result"]["total_sponsor_distributed"])
        total_cash_flow = sum(float(flow) for flow in solved.json()["result"]["underwriting"]["annual_cash_flows"])
        self.assertAlmostEqual(lp + sponsor, total_cash_flow, places=4)

    def test_waterfall_tier_order_is_selectable(self) -> None:
        pref_first = client.post(
            "/v1/waterfall",
            json={"underwriting": UNDERWRITING_REQUEST, "preferred_return_rate": "0.08", "sponsor_residual_share": "0.20"},
        )
        capital_first = client.post(
            "/v1/waterfall",
            json={
                "underwriting": UNDERWRITING_REQUEST,
                "preferred_return_rate": "0.08",
                "sponsor_residual_share": "0.20",
                "pay_preferred_return_before_capital": False,
            },
        )

        self.assertEqual(pref_first.status_code, 200)
        self.assertEqual(capital_first.status_code, 200)
        self.assertLess(
            float(pref_first.json()["total_sponsor_distributed"]),
            float(capital_first.json()["total_sponsor_distributed"]),
        )

    def test_exports_accept_report_context_and_describe_the_analysis_price(self) -> None:
        request = {
            **UNDERWRITING_REQUEST,
            "purchase_price": "1083333.33",
            "report": {
                "deal_name": "Maple Apartments",
                "version_label": "IC base case",
                "analysis_basis": "solved maximum supportable price",
                "target_levered_irr": "0.10",
                "minimum_dscr": "1.00",
                "maximum_purchase_price": "1083333.33",
                "recommended_bid_floor": "1029166.66",
                "recommended_bid_ceiling": "1083333.33",
                "binding_constraint": "minimum_dscr",
            },
        }

        workbook = client.post("/v1/exports/excel", json=request)
        memo = client.post("/v1/exports/memo", json=request)

        self.assertEqual(workbook.status_code, 200)
        summary = load_workbook(BytesIO(workbook.content), data_only=True)["Summary"]
        self.assertEqual(summary["A1"].value, "Maple Apartments")
        self.assertIn("Version: IC base case", summary["A2"].value)
        labels = {row[0].value: row[1].value for row in summary.iter_rows(min_row=3, max_col=2) if row[0].value}
        self.assertEqual(labels["Analysis basis"], "solved maximum supportable price")
        self.assertEqual(labels["Analysis purchase price"], 1083333.33)
        self.assertEqual(labels["Maximum supportable price"], 1083333.33)
        self.assertEqual(labels["Binding constraint"], "minimum dscr")
        self.assertEqual(memo.status_code, 200)
        self.assertTrue(memo.content.startswith(b"%PDF"))

    def test_year_one_growth_convention_is_selectable(self) -> None:
        grown = {
            **UNDERWRITING_REQUEST,
            "operating": {**UNDERWRITING_REQUEST["operating"], "rent_growth_rate": "0.10", "apply_growth_in_year_one": True},
        }
        current = {**grown, "operating": {**grown["operating"], "apply_growth_in_year_one": False}}

        grown_body = client.post("/v1/underwriting", json=grown).json()
        current_body = client.post("/v1/underwriting", json=current).json()

        self.assertEqual(grown_body["operating_projection"][0]["potential_gross_income"], "220000.00")
        self.assertEqual(current_body["operating_projection"][0]["potential_gross_income"], "200000")
        # Going-in cap is Year 1 NOI / price under either convention.
        self.assertEqual(current_body["going_in_cap_rate"], "0.13")
        self.assertAlmostEqual(float(grown_body["going_in_cap_rate"]), (220000 * 0.9 - 50000) / 1000000, places=12)


class DealPipelineApiTests(unittest.TestCase):
    """Deal list summaries, renaming, and deletion for the pipeline view."""

    def setUp(self) -> None:
        self.temp_directory = TemporaryDirectory()
        self.original_database = api_main.database
        database_path = Path(self.temp_directory.name) / "underwriting.db"
        api_main.database = Database(f"sqlite:///{database_path.as_posix()}")
        self.client_context = TestClient(api_main.app)
        self.client = self.client_context.__enter__()

    def tearDown(self) -> None:
        self.client_context.__exit__(None, None, None)
        api_main.database.dispose()
        api_main.database = self.original_database
        self.temp_directory.cleanup()

    def test_deal_list_summarises_latest_version_and_solve(self) -> None:
        empty = self.client.post("/v1/deals", json={"name": "Empty deal"}).json()
        deal = self.client.post("/v1/deals", json={"name": "Maple Apartments"}).json()
        version = self.client.post(
            f"/v1/deals/{deal['id']}/assumption-sets",
            json={"label": "Broker OM", "property_type": "Multifamily", "underwriting": UNDERWRITING_REQUEST},
        ).json()
        solved = self.client.post(
            f"/v1/assumption-sets/{version['id']}/maximum-price",
            json={"target_levered_irr": "0.10", "minimum_dscr": "1.00"},
        ).json()

        listing = self.client.get("/v1/deals")

        self.assertEqual(listing.status_code, 200)
        rows = {row["id"]: row for row in listing.json()}
        self.assertEqual(rows[empty["id"]]["version_count"], 0)
        self.assertIsNone(rows[empty["id"]]["latest_version"])
        self.assertIsNone(rows[empty["id"]]["latest_pricing"])
        maple = rows[deal["id"]]
        self.assertEqual(maple["version_count"], 1)
        self.assertEqual(maple["latest_version"]["label"], "Broker OM")
        self.assertEqual(maple["latest_pricing"]["maximum_purchase_price"], solved["result"]["maximum_purchase_price"])
        self.assertEqual(maple["latest_pricing"]["target_levered_irr"], "0.10")
        self.assertEqual(maple["latest_pricing"]["binding_constraint"], "minimum_dscr")
        self.assertEqual(maple["latest_pricing"]["risk_flag_count"], 2)
        # The solved deal was touched most recently, so it lists first.
        self.assertEqual(listing.json()[0]["id"], deal["id"])

    def test_deals_can_be_renamed_and_deleted_with_their_history(self) -> None:
        deal = self.client.post("/v1/deals", json={"name": "Temp"}).json()
        version = self.client.post(f"/v1/deals/{deal['id']}/assumption-sets", json={"underwriting": UNDERWRITING_REQUEST}).json()
        self.client.post(f"/v1/assumption-sets/{version['id']}/underwriting")

        renamed = self.client.patch(f"/v1/deals/{deal['id']}", json={"name": "  Oak Plaza  "})
        self.assertEqual(renamed.status_code, 200)
        self.assertEqual(renamed.json()["name"], "Oak Plaza")
        self.assertEqual(self.client.patch(f"/v1/deals/{deal['id']}", json={"name": "   "}).status_code, 422)

        deleted = self.client.delete(f"/v1/deals/{deal['id']}")
        self.assertEqual(deleted.status_code, 204)
        self.assertEqual(self.client.get(f"/v1/deals/{deal['id']}").status_code, 404)
        self.assertEqual(self.client.get(f"/v1/assumption-sets/{version['id']}").status_code, 404)
        self.assertEqual(self.client.delete(f"/v1/deals/{deal['id']}").status_code, 404)
