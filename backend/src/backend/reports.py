"""Downloadable underwriting workbook and investment-memo reports."""

from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle


HEADER_FILL = "166534"
RISK_FILL = "92400e"
MONEY = "$#,##0;($#,##0)"
PERCENT = "0.0%"
MULTIPLE = "0.00x"


def excel_cash_flow(result, assumptions=None, report=None):
    """Build a workbook with summary, cash flow, risk, and assumption sheets."""
    book = Workbook()
    summary = book.active
    summary.title = "Summary"
    cash = book.create_sheet("Cash Flow")
    risks = book.create_sheet("Risk Flags")
    assumption_sheet = book.create_sheet("Assumptions") if assumptions is not None else None

    summary.append(["Underwriting Summary", "Value"])
    summary_rows = _context_rows(report) + [
        ("Analysis purchase price", result.purchase_price, MONEY),
        ("Loan amount", result.loan_amount, MONEY),
        ("Closing costs", result.closing_costs, MONEY),
        ("Initial equity", result.initial_equity, MONEY),
        ("Going-in cap rate (Year 1 NOI / price)", result.going_in_cap_rate, PERCENT),
        ("Debt yield", result.debt_yield if result.debt_yield is not None else "N/A — all-cash", PERCENT),
        ("Levered IRR", result.returns.levered_irr, PERCENT),
        ("Equity multiple", result.returns.equity_multiple, MULTIPLE),
        ("Minimum DSCR", result.returns.minimum_dscr if result.returns.minimum_dscr is not None else "N/A — all-cash", MULTIPLE),
        ("Terminal NOI", result.terminal_noi, MONEY),
        ("Gross sale price", result.disposition.gross_sale_price, MONEY),
        ("Net sale proceeds to equity", result.disposition.net_sale_proceeds_to_equity, MONEY),
    ]
    for label, value, number_format in summary_rows:
        summary.append([label, value])
        if not isinstance(value, str):
            summary.cell(row=summary.max_row, column=2).number_format = number_format

    cash.append(["Year", "NOI", "Debt Service", "Operating Equity Cash Flow", "Sale Proceeds", "Total Equity Cash Flow", "Cash-on-Cash", "DSCR"])
    dscr_by_year = result.returns.dscr_by_year or (None,) * len(result.operating_projection)
    for operating, debt, operating_flow, flow, cash_on_cash, dscr in zip(
        result.operating_projection,
        result.debt_projection,
        result.annual_operating_cash_flows,
        result.annual_cash_flows,
        result.returns.cash_on_cash_by_year,
        dscr_by_year,
    ):
        cash.append([operating.year, operating.net_operating_income, debt.debt_service, operating_flow, flow - operating_flow, flow, cash_on_cash, dscr])
    for row in cash.iter_rows(min_row=2, min_col=2, max_col=6):
        for cell in row:
            cell.number_format = MONEY
    for cell in cash["G"][1:]:
        cell.number_format = PERCENT
    for cell in cash["H"][1:]:
        cell.number_format = MULTIPLE

    risks.append(["Code", "Severity", "Message", "Observed", "Threshold"])
    for flag in result.risk_flags:
        risks.append([flag.code, flag.severity, flag.message, flag.observed_value, flag.threshold])
    for row in risks.iter_rows(min_row=2, min_col=4, max_col=5):
        for cell in row:
            cell.number_format = "0.0000"

    if assumption_sheet is not None:
        assumption_sheet.append(["Underwriting assumptions", "Value"])
        for row in _assumption_rows(assumptions):
            assumption_sheet.append(row)

    for sheet in (summary, cash, risks, *([assumption_sheet] if assumption_sheet is not None else [])):
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        sheet.freeze_panes = "A2"
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 30 if column[0].column_letter == "A" else 22
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return output


def investment_memo(result, assumptions=None, report=None, detail=None):
    """Build a PDF memo whose language is tied to the stated return target."""
    output = BytesIO()
    doc = SimpleDocTemplate(output, pagesize=letter, rightMargin=42, leftMargin=42, topMargin=42, bottomMargin=42)
    styles = getSampleStyleSheet()
    styles["Heading2"].spaceBefore = 8
    styles["Heading2"].spaceAfter = 4
    body = styles["BodyText"]
    body.spaceBefore = 0
    body.spaceAfter = 0
    report = report or {}
    detail = detail or {}

    irr = float(result.returns.levered_irr)
    dscr = float(result.returns.minimum_dscr) if result.returns.minimum_dscr is not None else None
    target = report.get("target_levered_irr")
    minimum_dscr = report.get("minimum_dscr")
    title = "Investment Memo" if not report.get("deal_name") else f"Investment Memo — {report['deal_name']}"
    subtitle_parts = [part for part in (
        f"Version: {report['version_label']}" if report.get("version_label") else None,
        f"Effective {report['effective_date']}" if report.get("effective_date") else None,
        f"Confidence: {str(report['confidence_level']).title()}" if report.get("confidence_level") else None,
        f"Source: {report['source']}" if report.get("source") else None,
    ) if part]

    story = [Paragraph(title, styles["Title"])]
    if subtitle_parts:
        story.append(Paragraph(" · ".join(subtitle_parts), body))
    story.append(Spacer(1, 6))
    story.append(Paragraph(_pricing_paragraph(result, report), body))
    story.append(Spacer(1, 4))
    story.append(Paragraph(_return_opinion(irr, target, dscr, minimum_dscr), body))
    story.append(Spacer(1, 10))

    story.append(Paragraph("Return summary", styles["Heading2"]))
    rows = [
        ["Metric", "Result"],
        ["Analysis purchase price", _money(result.purchase_price)],
        ["Loan amount / initial equity", f"{_money(result.loan_amount)} / {_money(result.initial_equity)}"],
        ["Going-in cap rate (Year 1 NOI / price)", f"{float(result.going_in_cap_rate):.2%}"],
        ["Unlevered IRR / multiple", (f"{float(result.unlevered.irr):.1%} / " if result.unlevered.irr is not None else "n/a / ") + f"{float(result.unlevered.multiple):.2f}x"],
        ["Levered profit", _money(result.levered_profit)],
        ["Debt yield", f"{float(result.debt_yield):.2%}" if result.debt_yield is not None else "N/A — all-cash"],
        ["Levered IRR", f"{irr:.1%}" + (f" (target {float(target):.1%})" if target is not None else "")],
        ["Equity multiple", f"{float(result.returns.equity_multiple):.2f}x"],
        ["Minimum DSCR", (f"{dscr:.2f}x" + (f" (floor {float(minimum_dscr):.2f}x)" if minimum_dscr is not None else "")) if dscr is not None else "N/A — all-cash"],
        ["Terminal NOI", _money(result.terminal_noi)],
        ["Net sale proceeds to equity", _money(result.disposition.net_sale_proceeds_to_equity)],
    ]
    story += [_table(rows, [250, 190], HEADER_FILL), Spacer(1, 8)]
    property_detail = detail.get("property") or {}
    if property_detail.get("units") or property_detail.get("rentable_square_feet"):
        year_one = result.operating_projection[0]
        metric_rows = [["Metric", "Value"]]
        if property_detail.get("units"):
            units = float(property_detail["units"])
            metric_rows += [["Units", f"{units:,.0f}"], ["Price per unit", _money(float(result.purchase_price) / units)], ["Year 1 NOI per unit", _money(float(year_one.net_operating_income) / units)]]
        if property_detail.get("rentable_square_feet"):
            square_feet = float(property_detail["rentable_square_feet"])
            metric_rows += [["Rentable square feet", f"{square_feet:,.0f}"], ["Price per SF", f"${float(result.purchase_price) / square_feet:,.2f}"], ["Year 1 NOI per SF", f"${float(year_one.net_operating_income) / square_feet:,.2f}"]]
        metric_rows.append(["Year 1 expense ratio (OpEx / EGI)", f"{float(year_one.operating_expenses / year_one.effective_gross_income):.1%}" if year_one.effective_gross_income else "n/a"])
        story += [Paragraph("Property metrics", styles["Heading2"]), _table(metric_rows, [250, 190], HEADER_FILL, font_size=8.5), Spacer(1, 8)]

    story.append(Paragraph("Annual cash flow", styles["Heading2"]))
    cash_rows = [["Year", "NOI", "Debt service", "Recurring CF", "Cash-on-cash", "Total CF"]]
    for operating, debt, operating_flow, flow, cash_on_cash in zip(
        result.operating_projection, result.debt_projection, result.annual_operating_cash_flows, result.annual_cash_flows, result.returns.cash_on_cash_by_year
    ):
        cash_rows.append([str(operating.year), _money(operating.net_operating_income), _money(debt.debt_service), _money(operating_flow), f"{float(cash_on_cash):.1%}", _money(flow)])
    story += [_table(cash_rows, [40, 85, 85, 85, 70, 85], HEADER_FILL, font_size=8.5), Spacer(1, 8)]

    if assumptions is not None:
        story += [Paragraph("Key assumptions", styles["Heading2"]), _table([["Assumption", "Value"]] + _assumption_rows(assumptions), [250, 190], HEADER_FILL, font_size=8.5), Spacer(1, 8)]

    story.append(Paragraph("Risk review", styles["Heading2"]))
    if result.risk_flags:
        risk_rows = [["Severity", "Flag"]] + [[flag.severity.title(), flag.message] for flag in result.risk_flags]
        story += [_table(risk_rows, [90, 350], RISK_FILL), Spacer(1, 8)]
    else:
        story += [Paragraph("No rules-based risk flags were triggered by the modeled assumptions.", body), Spacer(1, 12)]

    story += [
        Paragraph("Basis and limitations", styles["Heading2"]),
        Paragraph(
            "Figures are model outputs for the stated assumptions, not a valuation or investment advice. "
            "The sale is capitalized on forward (year after hold) NOI, debt amortizes monthly, and cash-on-cash excludes sale proceeds. "
            "Review the exit-cap, rent-growth, and interest-rate sensitivities before relying on the base case.",
            body,
        ),
    ]
    doc.build(story)
    output.seek(0)
    return output


def _pricing_paragraph(result, report) -> str:
    maximum = report.get("maximum_purchase_price")
    if maximum is None:
        basis = report.get("analysis_basis") or "the entered purchase price"
        return f"This memo analyses the deal at {_money(result.purchase_price)} ({basis})."
    floor, ceiling = report.get("recommended_bid_floor"), report.get("recommended_bid_ceiling")
    binding = {"minimum_dscr": "the minimum DSCR floor", "target_levered_irr": "the target levered IRR"}.get(report.get("binding_constraint"), None)
    text = f"The reverse solver's maximum supportable price is {_money(maximum)}"
    if binding:
        text += f", limited by {binding}"
    text += f". This memo analyses the deal at {_money(result.purchase_price)}"
    text += f" ({report['analysis_basis']})." if report.get("analysis_basis") else "."
    if floor is not None and ceiling is not None:
        text += f" The recommended bid range is {_money(floor)} to {_money(ceiling)} (95% of the maximum price up to the maximum itself)."
    return text


def _return_opinion(irr: float, target, dscr, minimum_dscr) -> str:
    if target is not None:
        gap = irr - float(target)
        if gap >= 0:
            opinion = f"The modeled levered IRR of {irr:.1%} meets the {float(target):.1%} target by {gap:.1%}."
        else:
            opinion = f"The modeled levered IRR of {irr:.1%} falls {abs(gap):.1%} short of the {float(target):.1%} target."
    else:
        opinion = f"The modeled levered IRR is {irr:.1%}; no return target was supplied for comparison."
    if dscr is None:
        coverage = "No debt service is modeled, so DSCR and debt yield are not applicable."
    elif minimum_dscr is not None:
        coverage = (
            f"Minimum DSCR of {dscr:.2f}x satisfies the {float(minimum_dscr):.2f}x floor."
            if dscr >= float(minimum_dscr)
            else f"Minimum DSCR of {dscr:.2f}x is below the {float(minimum_dscr):.2f}x floor."
        )
    else:
        coverage = f"Minimum DSCR is {dscr:.2f}x." + (" Debt coverage is thin and warrants caution." if dscr < 1.20 else "")
    return f"{opinion} {coverage}"


def _context_rows(report):
    if not report:
        return []
    rows = []
    for key, label in (("deal_name", "Deal"), ("version_label", "Version"), ("source", "Source"), ("effective_date", "Effective date"), ("confidence_level", "Confidence"), ("analysis_basis", "Analysis basis")):
        if report.get(key):
            rows.append((label, str(report[key]), None))
    for key, label, number_format in (
        ("target_levered_irr", "Target levered IRR", PERCENT),
        ("minimum_dscr", "Minimum DSCR floor", MULTIPLE),
        ("maximum_purchase_price", "Maximum supportable price", MONEY),
        ("recommended_bid_floor", "Recommended bid floor", MONEY),
        ("recommended_bid_ceiling", "Recommended bid ceiling", MONEY),
    ):
        if report.get(key) is not None:
            rows.append((label, report[key], number_format))
    if report.get("binding_constraint"):
        rows.append(("Binding constraint", report["binding_constraint"].replace("_", " "), None))
    return rows


def _table(rows, widths, header_fill, font_size=9):
    table = Table(rows, colWidths=widths)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{header_fill}")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
                ("PADDING", (0, 0), (-1, -1), 3),
                ("FONTSIZE", (0, 0), (-1, -1), font_size),
                ("ALIGN", (1, 1), (-1, -1), "RIGHT"),
            ]
        )
    )
    return table


def _money(value) -> str:
    amount = float(value)
    return f"(${abs(amount):,.0f})" if amount < 0 else f"${amount:,.0f}"


def _assumption_rows(assumptions):
    operating = assumptions.operating
    schedule = operating.annual_rent_growth_rates
    schedule_display = ", ".join(f"{float(rate):.1%}" for rate in schedule) if schedule else "Not supplied (base rate used)"
    return [
        ["Purchase price", _money(assumptions.purchase_price)],
        ["Closing costs", f"{float(assumptions.closing_cost_rate):.1%}"],
        ["Loan-to-value", f"{float(assumptions.loan_to_value):.1%}"],
        ["Annual interest rate", f"{float(assumptions.annual_interest_rate):.2%}"],
        ["Amortization", f"{assumptions.amortization_years} years"],
        ["Exit cap rate", f"{float(assumptions.exit_cap_rate):.2%}"],
        ["Selling costs", f"{float(assumptions.selling_cost_rate):.1%}"],
        ["Current potential gross income", _money(operating.current_potential_gross_income)],
        ["Current operating expenses", _money(operating.current_operating_expenses)],
        ["Vacancy", f"{float(operating.vacancy_rate):.1%}"],
        ["Base rent growth", f"{float(operating.rent_growth_rate):.1%}"],
        ["Annual rent-growth schedule", schedule_display],
        ["Expense growth", f"{float(operating.expense_growth_rate):.1%}"],
        ["Growth convention", "Growth applied before Year 1" if operating.apply_growth_in_year_one else "Year 1 = current figures; growth from Year 2"],
        ["Hold period", f"{operating.hold_period_years} years"],
    ]
