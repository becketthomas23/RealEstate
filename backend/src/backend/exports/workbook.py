"""Build the analyst-grade Excel underwriting workbook.

Every number in the workbook is produced by the audited calculation engine.
Subtotals, coverage ratios, and return metrics are written as live Excel
formulas over those values so an analyst can trace each line and the sheet
recalculates if a cell is overridden; Excel is told to recalculate on open.
"""

from __future__ import annotations

from dataclasses import replace
from decimal import Decimal
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from backend.core.validation import ZERO
from backend.solvers.pricing import MaximumPriceAssumptions, solve_maximum_purchase_price
from backend.workflows.sensitivity import SensitivityAssumptions, calculate_one_way_sensitivities
from backend.workflows.underwriting import UnderwritingAssumptions, UnderwritingResult, calculate_underwriting
from backend.workflows.underwriting.scenarios import apply_scenario_values


GREEN = "166534"
GREEN_SOFT = "E6F2EA"
INK = "16201B"
MUTED = "7A847E"
LINE = "D9DED8"
AMBER = "92400E"
AMBER_SOFT = "FFF4E5"
RED_SOFT = "FDECEC"
MONEY = "$#,##0;($#,##0);-"
MONEY_CENTS = "$#,##0.00;($#,##0.00);-"
PERCENT = "0.00%"
PERCENT_1 = "0.0%"
MULTIPLE = "0.00x"
INTEGER = "#,##0"
RATIO = "0.00"

TITLE_FONT = Font(bold=True, size=16, color=INK)
SUBTITLE_FONT = Font(size=10, color=MUTED)
SECTION_FONT = Font(bold=True, size=11, color="FFFFFF")
SECTION_FILL = PatternFill("solid", fgColor=GREEN)
HEADER_FONT = Font(bold=True, color=INK)
HEADER_FILL = PatternFill("solid", fgColor=GREEN_SOFT)
TOTAL_FONT = Font(bold=True, color=INK)
LABEL_FONT = Font(color=INK)
NOTE_FONT = Font(italic=True, size=9, color=MUTED)
THIN = Side(style="thin", color=LINE)
TOP_BORDER = Border(top=THIN)
BOTTOM_BORDER = Border(bottom=THIN)


def build_underwriting_workbook(
    result: UnderwritingResult,
    assumptions: UnderwritingAssumptions,
    *,
    report: dict[str, Any] | None = None,
    detail: dict[str, Any] | None = None,
    simulation=None,
    include_sensitivity: bool = True,
) -> BytesIO:
    report = report or {}
    detail = detail or {}
    book = Workbook()
    book.calculation.fullCalcOnLoad = True
    _summary_sheet(book.active, result, assumptions, report, detail)
    _pro_forma_sheet(book.create_sheet("Pro Forma"), result, assumptions, report)
    _monthly_sheet(book.create_sheet("Monthly Cash Flow"), result, assumptions)
    _debt_sheet(book.create_sheet("Debt"), result, assumptions)
    _returns_sheet(book.create_sheet("Sale & Returns"), result, assumptions, report)
    if include_sensitivity:
        _sensitivity_sheet(book.create_sheet("Sensitivity"), result, assumptions, report)
    if simulation is not None:
        _risk_sheet(book.create_sheet("Risk Simulation"), simulation, report)
    if detail.get("rent_roll") or detail.get("operating_expenses") or detail.get("other_income") or detail.get("property"):
        _detail_sheet(book.create_sheet("Rent Roll & Expenses"), result, assumptions, detail)
    _flags_sheet(book.create_sheet("Risk Flags"), result)
    _assumptions_sheet(book.create_sheet("Assumptions"), assumptions, report, detail)
    for sheet in book.worksheets:
        sheet.sheet_view.showGridLines = False
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
    output = BytesIO()
    book.save(output)
    output.seek(0)
    return output


# ================================================================ Summary


def _summary_sheet(sheet: Worksheet, result: UnderwritingResult, assumptions: UnderwritingAssumptions, report: dict, detail: dict) -> None:
    sheet.title = "Summary"
    _widths(sheet, {"A": 34, "B": 18, "C": 4, "D": 34, "E": 18, "F": 4, "G": 30, "H": 18})
    hold = assumptions.operating.hold_period_years
    property_detail = detail.get("property") or {}
    deal_name = report.get("deal_name") or property_detail.get("name") or "Underwriting summary"
    row = _title(sheet, deal_name, _subtitle(report, assumptions, result))

    # ---- Pricing conclusion (left) and returns (middle) and capitalization (right)
    row = _section(sheet, row, "Pricing conclusion", "A", "B")
    _section(sheet, row - 1, "Returns", "D", "E")
    _section(sheet, row - 1, "Sources and uses", "G", "H")
    returns = result.returns
    target = report.get("target_levered_irr")
    floor = report.get("minimum_dscr")
    left = [
        ("Analysis purchase price", result.purchase_price, MONEY),
        ("Analysis basis", report.get("analysis_basis") or "entered purchase price", None),
        ("Maximum supportable price", report.get("maximum_purchase_price"), MONEY),
        ("Recommended bid floor", report.get("recommended_bid_floor"), MONEY),
        ("Recommended bid ceiling", report.get("recommended_bid_ceiling"), MONEY),
        ("Binding constraint", (report.get("binding_constraint") or "").replace("_", " ") or None, None),
        ("Target levered IRR", target, PERCENT),
        ("Minimum DSCR floor", floor, MULTIPLE),
        ("Going-in cap rate (Year 1 NOI / price)", result.going_in_cap_rate, PERCENT),
        ("Exit cap rate", assumptions.exit_cap_rate, PERCENT),
        ("Cap rate spread (exit less going-in)", assumptions.exit_cap_rate - result.going_in_cap_rate, PERCENT),
        ("Hold period", f"{hold} years", None),
    ]
    middle = [
        ("Levered IRR", returns.levered_irr, PERCENT),
        ("Unlevered IRR", result.unlevered.irr, PERCENT),
        ("Levered equity multiple", returns.equity_multiple, MULTIPLE),
        ("Unlevered multiple", result.unlevered.multiple, MULTIPLE),
        ("Levered profit", result.levered_profit, MONEY),
        ("Unlevered profit", result.unlevered.profit, MONEY),
        ("Year 1 cash-on-cash", returns.cash_on_cash_by_year[0], PERCENT),
        ("Average cash-on-cash", sum(returns.cash_on_cash_by_year, ZERO) / len(returns.cash_on_cash_by_year), PERCENT),
        ("Minimum DSCR", returns.minimum_dscr if returns.minimum_dscr is not None else "N/A — all-cash", MULTIPLE),
        ("Year 1 debt yield", result.debt_yield if result.debt_yield is not None else "N/A — all-cash", PERCENT),
        ("Net sale proceeds to equity", result.disposition.net_sale_proceeds_to_equity, MONEY),
        ("Share of profit from sale", _share_of_profit_from_sale(result), PERCENT),
    ]
    total_cost = result.unlevered.total_cost
    right = [
        ("Uses", None, None),
        ("Purchase price", result.purchase_price, MONEY),
        ("Closing costs", result.closing_costs, MONEY),
        ("Total uses", total_cost, MONEY),
        ("Sources", None, None),
        ("Senior loan", result.loan_amount, MONEY),
        ("Equity", result.initial_equity, MONEY),
        ("Total sources", result.loan_amount + result.initial_equity, MONEY),
        ("Loan-to-value", assumptions.loan_to_value, PERCENT),
        ("Loan-to-cost", result.loan_amount / total_cost if total_cost else ZERO, PERCENT),
        ("Equity share of total cost", result.initial_equity / total_cost if total_cost else ZERO, PERCENT),
    ]
    end = row
    for offset, (label, value, fmt) in enumerate(left):
        _kv(sheet, row + offset, "A", "B", label, value, fmt)
    for offset, (label, value, fmt) in enumerate(middle):
        _kv(sheet, row + offset, "D", "E", label, value, fmt)
    for offset, (label, value, fmt) in enumerate(right):
        bold = label in {"Uses", "Sources", "Total uses", "Total sources"}
        _kv(sheet, row + offset, "G", "H", label, value, fmt, bold=bold)
    end = row + max(len(left), len(middle), len(right)) + 1

    # ---- Per-unit and per-square-foot metrics
    units = property_detail.get("units")
    square_feet = property_detail.get("rentable_square_feet")
    if units or square_feet:
        end = _section(sheet, end, "Per-unit and per-square-foot metrics", "A", "H")
        year_one = result.operating_projection[0]
        metrics = []
        if units:
            units_d = Decimal(str(units))
            metrics += [
                ("Price per unit", result.purchase_price / units_d, MONEY),
                ("Year 1 NOI per unit", year_one.net_operating_income / units_d, MONEY),
                ("Year 1 operating expenses per unit", year_one.operating_expenses / units_d, MONEY),
                ("Year 1 gross income per unit per month", year_one.potential_gross_income / units_d / 12, MONEY),
                ("Loan per unit", result.loan_amount / units_d, MONEY),
            ]
        if square_feet:
            sf = Decimal(str(square_feet))
            metrics += [
                ("Price per rentable SF", result.purchase_price / sf, MONEY_CENTS),
                ("Year 1 NOI per SF", year_one.net_operating_income / sf, MONEY_CENTS),
                ("Year 1 operating expenses per SF", year_one.operating_expenses / sf, MONEY_CENTS),
                ("Year 1 gross income per SF", year_one.potential_gross_income / sf, MONEY_CENTS),
            ]
        metrics += [
            ("Year 1 operating expense ratio (OpEx / EGI)", year_one.operating_expenses / year_one.effective_gross_income if year_one.effective_gross_income else ZERO, PERCENT),
            ("Year 1 NOI margin (NOI / EGI)", year_one.net_operating_income / year_one.effective_gross_income if year_one.effective_gross_income else ZERO, PERCENT),
            ("Break-even occupancy ((OpEx + debt service) / PGI)", (year_one.operating_expenses + result.debt_projection[0].debt_service) / year_one.potential_gross_income, PERCENT),
        ]
        columns = (("A", "B"), ("D", "E"), ("G", "H"))
        for index, (label, value, fmt) in enumerate(metrics):
            label_column, value_column = columns[index % 3]
            _kv(sheet, end + index // 3, label_column, value_column, label, value, fmt)
        end += (len(metrics) + 2) // 3 + 1

    # ---- Annual snapshot
    end = _section(sheet, end, "Annual snapshot", "A", "H")
    headers = ["Year", "NOI", "Debt service", "Cash flow after debt", "Cash-on-cash", "DSCR", "Debt yield", "Levered cash flow"]
    for index, header in enumerate(headers):
        _header_cell(sheet.cell(row=end, column=index + 1, value=header))
    end += 1
    for index, year in enumerate(result.operating_projection):
        debt = result.debt_projection[index]
        values = [
            year.year,
            year.net_operating_income,
            debt.debt_service,
            result.annual_operating_cash_flows[index],
            returns.cash_on_cash_by_year[index],
            returns.dscr_by_year[index] if returns.dscr_by_year else None,
            result.annual_debt_yield[index] if result.annual_debt_yield else None,
            result.annual_cash_flows[index],
        ]
        formats = [INTEGER, MONEY, MONEY, MONEY, PERCENT_1, MULTIPLE, PERCENT_1, MONEY]
        for column, (value, fmt) in enumerate(zip(values, formats), start=1):
            cell = sheet.cell(row=end, column=column, value=_num(value) if value is not None else "N/A")
            cell.number_format = fmt
            cell.alignment = Alignment(horizontal="right")
        end += 1
    end += 1

    # ---- Risk flags
    end = _section(sheet, end, "Risk flags", "A", "H")
    if result.risk_flags:
        for flag in result.risk_flags:
            sheet.cell(row=end, column=1, value=flag.severity.title()).font = Font(bold=True, color=AMBER if flag.severity != "critical" else "B91C1C")
            sheet.cell(row=end, column=2, value=flag.message)
            sheet.merge_cells(start_row=end, start_column=2, end_row=end, end_column=8)
            end += 1
    else:
        sheet.cell(row=end, column=1, value="No rules-based risk flags were triggered at the analysis price.").font = LABEL_FONT
        end += 1
    end += 1
    end = _section(sheet, end, "Basis and conventions", "A", "H")
    for note in _convention_notes(assumptions):
        cell = sheet.cell(row=end, column=1, value=note)
        cell.font = NOTE_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.merge_cells(start_row=end, start_column=1, end_row=end, end_column=8)
        sheet.row_dimensions[end].height = 28
        end += 1
    sheet.freeze_panes = "A4"


def _subtitle(report: dict, assumptions: UnderwritingAssumptions, result: UnderwritingResult) -> str:
    parts = []
    if report.get("version_label"):
        parts.append(f"Version: {report['version_label']}")
    if report.get("effective_date"):
        parts.append(f"Effective {report['effective_date']}")
    if report.get("confidence_level"):
        parts.append(f"Confidence: {str(report['confidence_level']).title()}")
    if report.get("source"):
        parts.append(f"Source: {report['source']}")
    parts.append(f"{assumptions.operating.hold_period_years}-year hold")
    return " · ".join(parts)


def _share_of_profit_from_sale(result: UnderwritingResult) -> Decimal:
    if result.levered_profit <= ZERO:
        return ZERO
    sale_gain = result.disposition.net_sale_proceeds_to_equity - result.initial_equity
    return max(min(sale_gain / result.levered_profit, Decimal("1")), ZERO)


# ============================================================== Pro Forma


def _pro_forma_sheet(sheet: Worksheet, result: UnderwritingResult, assumptions: UnderwritingAssumptions, report: dict) -> None:
    hold = assumptions.operating.hold_period_years
    years = list(result.operating_projection)
    terminal = replace(years[-1], year=hold + 1)  # placeholder for the forward column
    sheet.column_dimensions["A"].width = 40
    sheet.column_dimensions["B"].width = 16
    for index in range(hold + 1):
        sheet.column_dimensions[get_column_letter(index + 3)].width = 15
    row = _title(sheet, "Annual pro forma", "Values from the calculation engine; subtotals, ratios and returns are live formulas. Year 0 is closing.")
    header_row = row
    sheet.cell(row=row, column=1, value="Line item")
    sheet.cell(row=row, column=2, value="Year 0")
    for index, year in enumerate(years):
        sheet.cell(row=row, column=index + 3, value=f"Year {year.year}")
    sheet.cell(row=row, column=hold + 3, value=f"Year {hold + 1} (forward)")
    for cell in sheet[row]:
        _header_cell(cell)
    row += 1
    first_year_column = 3
    last_year_column = hold + 2
    forward_column = hold + 3
    col = get_column_letter

    def line(label: str, values: list, fmt: str, *, bold=False, year0=None, forward=None, formula=None, note=False):
        nonlocal row
        label_cell = sheet.cell(row=row, column=1, value=label)
        label_cell.font = TOTAL_FONT if bold else (NOTE_FONT if note else LABEL_FONT)
        if year0 is not None:
            cell = sheet.cell(row=row, column=2, value=year0)
            cell.number_format = fmt
            cell.font = TOTAL_FONT if bold else LABEL_FONT
        for index in range(hold):
            column = first_year_column + index
            value = formula(col(column)) if formula else _num(values[index])
            cell = sheet.cell(row=row, column=column, value=value)
            cell.number_format = fmt
            if bold:
                cell.font = TOTAL_FONT
                cell.border = TOP_BORDER
        if forward is not None:
            cell = sheet.cell(row=row, column=forward_column, value=forward(col(forward_column)) if callable(forward) else _num(forward))
            cell.number_format = fmt
            cell.font = TOTAL_FONT if bold else LABEL_FONT
        current = row
        row += 1
        return current

    row = _section(sheet, row, "Operations", "A", col(forward_column))
    terminal_year = _forward_year(assumptions)
    r_pgi = line("Potential gross income", [y.potential_gross_income for y in years], MONEY, forward=terminal_year.potential_gross_income)
    r_vac = line("Less: vacancy and credit loss", [-y.vacancy_loss for y in years], MONEY, forward=-terminal_year.vacancy_loss)
    r_egi = line("Effective gross income", [], MONEY, bold=True, formula=lambda c: f"={c}{r_pgi}+{c}{r_vac}", forward=lambda c: f"={c}{r_pgi}+{c}{r_vac}")
    r_opex = line("Less: operating expenses", [-y.operating_expenses for y in years], MONEY, forward=-terminal_year.operating_expenses)
    r_noi = line("Net operating income", [], MONEY, bold=True, formula=lambda c: f"={c}{r_egi}+{c}{r_opex}", forward=lambda c: f"={c}{r_egi}+{c}{r_opex}")
    line("NOI margin (NOI / EGI)", [], PERCENT_1, formula=lambda c: f"=IF({c}{r_egi}=0,0,{c}{r_noi}/{c}{r_egi})", note=True)
    growth_cells = []
    for index in range(hold):
        column = first_year_column + index
        growth_cells.append(f"=IF({col(column - 1)}{r_noi}=0,\"\",{col(column)}{r_noi}/{col(column - 1)}{r_noi}-1)" if index > 0 else "")
    sheet.cell(row=row, column=1, value="NOI growth (year over year)").font = NOTE_FONT
    for index, formula in enumerate(growth_cells):
        cell = sheet.cell(row=row, column=first_year_column + index, value=formula if formula else None)
        cell.number_format = PERCENT_1
    row += 2

    row = _section(sheet, row, "Debt service", "A", col(forward_column))
    r_int = line("Interest", [-d.interest for d in result.debt_projection], MONEY)
    r_prin = line("Principal", [-d.principal for d in result.debt_projection], MONEY)
    r_ds = line("Total debt service", [], MONEY, bold=True, formula=lambda c: f"={c}{r_int}+{c}{r_prin}")
    r_cf = line("Cash flow after debt service", [], MONEY, bold=True, formula=lambda c: f"={c}{r_noi}+{c}{r_ds}")
    line("Debt service coverage (NOI / debt service)", [], MULTIPLE, formula=lambda c: f"=IF({c}{r_ds}=0,\"N/A\",{c}{r_noi}/-{c}{r_ds})", note=True)
    line("Debt yield (NOI / loan amount)", [], PERCENT_1, formula=lambda c: (f"={c}{r_noi}/$B${0}" if result.loan_amount > ZERO else "=\"N/A\""), note=True)
    r_debt_yield = row - 1
    line("Ending loan balance", [d.ending_balance for d in result.debt_projection], MONEY, year0=_num(result.loan_amount), note=True)
    row += 1

    row = _section(sheet, row, "Sale", "A", col(forward_column))
    sale_values = lambda value: [value if index == hold - 1 else ZERO for index in range(hold)]  # noqa: E731
    r_gross = line("Gross sale price (forward NOI / exit cap)", sale_values(result.disposition.gross_sale_price), MONEY)
    r_sell = line("Less: selling costs", sale_values(-result.disposition.selling_costs), MONEY)
    r_payoff = line("Less: loan payoff", sale_values(-result.disposition.loan_payoff), MONEY)
    r_net = line("Net sale proceeds to equity", [], MONEY, bold=True, formula=lambda c: f"={c}{r_gross}+{c}{r_sell}+{c}{r_payoff}")
    row += 1

    row = _section(sheet, row, "Equity cash flow", "A", col(forward_column))
    r_equity = line("Equity investment (price + closing costs − loan)", [ZERO] * hold, MONEY, year0=_num(-result.initial_equity))
    r_lev = line("Levered cash flow", [], MONEY, bold=True, year0=f"=B{r_equity}", formula=lambda c: f"={c}{r_cf}+{c}{r_net}")
    r_cum = row
    sheet.cell(row=row, column=1, value="Cumulative levered cash flow").font = NOTE_FONT
    sheet.cell(row=row, column=2, value=f"=B{r_lev}").number_format = MONEY
    for index in range(hold):
        column = first_year_column + index
        cell = sheet.cell(row=row, column=column, value=f"={col(column - 1)}{r_cum}+{col(column)}{r_lev}")
        cell.number_format = MONEY
    row += 1
    line("Cash-on-cash (cash flow after debt / equity)", [], PERCENT_1, formula=lambda c: f"=IF($B${r_equity}=0,0,{c}{r_cf}/-$B${r_equity})", note=True)
    row += 1

    row = _section(sheet, row, "Unlevered cash flow", "A", col(forward_column))
    r_cost = line("Total cost (price + closing costs)", [ZERO] * hold, MONEY, year0=_num(-result.unlevered.total_cost))
    r_unlev = line("Unlevered cash flow (NOI + net sale before debt)", [], MONEY, bold=True, year0=f"=B{r_cost}", formula=lambda c: f"={c}{r_noi}+{c}{r_gross}+{c}{r_sell}")
    row += 1

    row = _section(sheet, row, "Returns (live formulas over the rows above)", "A", col(forward_column))
    first, last = col(2), col(last_year_column)
    for label, formula, fmt, engine in (
        ("Levered IRR", f"=IRR({first}{r_lev}:{last}{r_lev})", PERCENT, result.returns.levered_irr),
        ("Levered equity multiple", f"=SUMIF({col(first_year_column)}{r_lev}:{last}{r_lev},\">0\")/-(B{r_lev}+SUMIF({col(first_year_column)}{r_lev}:{last}{r_lev},\"<0\"))", MULTIPLE, result.returns.equity_multiple),
        ("Levered profit", f"=SUM({first}{r_lev}:{last}{r_lev})", MONEY, result.levered_profit),
        ("Unlevered IRR", f"=IRR({first}{r_unlev}:{last}{r_unlev})", PERCENT, result.unlevered.irr),
        ("Unlevered multiple", f"=SUMIF({col(first_year_column)}{r_unlev}:{last}{r_unlev},\">0\")/-(B{r_unlev}+SUMIF({col(first_year_column)}{r_unlev}:{last}{r_unlev},\"<0\"))", MULTIPLE, result.unlevered.multiple),
        ("Minimum DSCR", f"=IF(B{r_ds}=\"\",MIN({col(first_year_column)}{r_ds + 2}:{last}{r_ds + 2}),MIN({col(first_year_column)}{r_ds + 2}:{last}{r_ds + 2}))" if result.loan_amount > ZERO else "=\"N/A\"", MULTIPLE, result.returns.minimum_dscr),
    ):
        sheet.cell(row=row, column=1, value=label).font = TOTAL_FONT
        cell = sheet.cell(row=row, column=2, value=formula)
        cell.number_format = fmt
        cell.font = TOTAL_FONT
        check = sheet.cell(row=row, column=3, value=_num(engine) if engine is not None else "N/A")
        check.number_format = fmt
        check.font = NOTE_FONT
        sheet.cell(row=row, column=4, value="← engine value for reference").font = NOTE_FONT
        row += 1
    # Fix the debt-yield formula to point at the loan amount cell written on the ending-balance line.
    loan_cell = f"$B${r_debt_yield + 1}"
    for index in range(hold):
        cell = sheet.cell(row=r_debt_yield, column=first_year_column + index)
        if result.loan_amount > ZERO:
            cell.value = f"={col(first_year_column + index)}{r_noi}/{loan_cell}"
    sheet.freeze_panes = f"C{header_row + 1}"
    _negatives_red(sheet, f"B{header_row + 1}:{col(forward_column)}{row}")


def _forward_year(assumptions: UnderwritingAssumptions):
    from backend.engines.operating import project_operating_year

    return project_operating_year(assumptions.operating, assumptions.operating.hold_period_years + 1)


# ================================================================ Monthly


def _monthly_sheet(sheet: Worksheet, result: UnderwritingResult, assumptions: UnderwritingAssumptions) -> None:
    row = _title(sheet, "Monthly cash flow", "Annual operating amounts allocated evenly across each year; debt service from the monthly amortization schedule.")
    headers = ["Month", "Year", "Potential gross income", "Vacancy and credit loss", "Effective gross income", "Operating expenses", "NOI", "Interest", "Principal", "Debt service", "Cash flow after debt", "Ending loan balance"]
    for index, header in enumerate(headers, start=1):
        _header_cell(sheet.cell(row=row, column=index, value=header))
    _widths(sheet, {get_column_letter(i): 16 for i in range(1, len(headers) + 1)})
    sheet.column_dimensions["A"].width = 8
    sheet.column_dimensions["B"].width = 6
    row += 1
    from backend.engines.debt import DebtAssumptions, project_monthly_debt

    monthly_debt = (
        project_monthly_debt(
            DebtAssumptions(result.loan_amount, assumptions.annual_interest_rate, assumptions.amortization_years, assumptions.operating.hold_period_years)
        )
        if result.loan_amount > ZERO
        else None
    )
    first = row
    for index, month in enumerate(result.monthly_operating_projection):
        debt = monthly_debt[index] if monthly_debt else None
        values = [
            month.month,
            month.year,
            month.potential_gross_income,
            -month.vacancy_loss,
            None,
            -month.operating_expenses,
            None,
            -(debt.interest if debt else ZERO),
            -(debt.principal if debt else ZERO),
            None,
            None,
            debt.ending_balance if debt else ZERO,
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column)
            if column == 5:
                cell.value = f"=C{row}+D{row}"
            elif column == 7:
                cell.value = f"=E{row}+F{row}"
            elif column == 10:
                cell.value = f"=H{row}+I{row}"
            elif column == 11:
                cell.value = f"=G{row}+J{row}"
            else:
                cell.value = _num(value)
            cell.number_format = INTEGER if column <= 2 else MONEY
        row += 1
    sheet.cell(row=row, column=1, value="Total").font = TOTAL_FONT
    for column in range(3, 12):
        letter = get_column_letter(column)
        cell = sheet.cell(row=row, column=column, value=f"=SUM({letter}{first}:{letter}{row - 1})")
        cell.number_format = MONEY
        cell.font = TOTAL_FONT
        cell.border = TOP_BORDER
    sheet.freeze_panes = "C4"


# =================================================================== Debt


def _debt_sheet(sheet: Worksheet, result: UnderwritingResult, assumptions: UnderwritingAssumptions) -> None:
    _widths(sheet, {"A": 36, "B": 18, "C": 4, "D": 10, "E": 16, "F": 16, "G": 16, "H": 16, "I": 16, "J": 16})
    row = _title(sheet, "Debt", "Fixed-rate senior loan, monthly principal and interest, repaid from sale proceeds.")
    if result.loan_amount <= ZERO:
        sheet.cell(row=row, column=1, value="All-cash acquisition: no debt is modeled.").font = LABEL_FONT
        return
    from backend.engines.debt import DebtAssumptions, project_monthly_debt

    monthly = project_monthly_debt(
        DebtAssumptions(result.loan_amount, assumptions.annual_interest_rate, assumptions.amortization_years, assumptions.operating.hold_period_years)
    )
    row = _section(sheet, row, "Loan summary", "A", "B")
    hold = assumptions.operating.hold_period_years
    summary = [
        ("Loan amount", result.loan_amount, MONEY),
        ("Loan-to-value", assumptions.loan_to_value, PERCENT),
        ("Loan-to-cost", result.loan_amount / result.unlevered.total_cost, PERCENT),
        ("Interest rate (annual, monthly compounding)", assumptions.annual_interest_rate, PERCENT),
        ("Amortization", f"{assumptions.amortization_years} years", None),
        ("Monthly payment", monthly[0].debt_service, MONEY_CENTS),
        ("Annual debt service", result.debt_projection[0].debt_service, MONEY),
        ("Loan constant (annual debt service / loan)", result.debt_projection[0].debt_service / result.loan_amount, PERCENT),
        ("Year 1 debt yield", result.debt_yield, PERCENT),
        ("Minimum DSCR over hold", result.returns.minimum_dscr, MULTIPLE),
        ("Balance at sale", result.disposition.loan_payoff, MONEY),
        ("Principal repaid over hold", result.loan_amount - result.disposition.loan_payoff, MONEY),
        ("Share of loan repaid", (result.loan_amount - result.disposition.loan_payoff) / result.loan_amount, PERCENT),
        ("Total interest over hold", sum((year.interest for year in result.debt_projection), ZERO), MONEY),
    ]
    for offset, (label, value, fmt) in enumerate(summary):
        _kv(sheet, row + offset, "A", "B", label, value, fmt)

    # Annual schedule to the right
    annual_row = row - 1
    _section(sheet, annual_row, "Annual schedule", "D", "J")
    headers = ["Year", "Beginning balance", "Debt service", "Interest", "Principal", "Ending balance", "DSCR"]
    for index, header in enumerate(headers):
        _header_cell(sheet.cell(row=annual_row + 1, column=4 + index, value=header))
    for index, year in enumerate(result.debt_projection):
        r = annual_row + 2 + index
        values = [year.year, year.beginning_balance, year.debt_service, year.interest, year.principal, year.ending_balance, result.returns.dscr_by_year[index]]
        formats = [INTEGER, MONEY, MONEY, MONEY, MONEY, MONEY, MULTIPLE]
        for offset, (value, fmt) in enumerate(zip(values, formats)):
            cell = sheet.cell(row=r, column=4 + offset, value=_num(value))
            cell.number_format = fmt
    row = max(row + len(summary), annual_row + 2 + hold) + 1

    row = _section(sheet, row, "Monthly amortization schedule", "A", "J")
    headers = ["Month", "Year", "Beginning balance", "Payment", "Interest", "Principal", "Ending balance"]
    for index, header in enumerate(headers, start=1):
        _header_cell(sheet.cell(row=row, column=index, value=header))
    row += 1
    for month in monthly:
        values = [month.month, month.year, month.beginning_balance, month.debt_service, month.interest, month.principal, month.ending_balance]
        formats = [INTEGER, INTEGER, MONEY_CENTS, MONEY_CENTS, MONEY_CENTS, MONEY_CENTS, MONEY_CENTS]
        for column, (value, fmt) in enumerate(zip(values, formats), start=1):
            cell = sheet.cell(row=row, column=column, value=_num(value))
            cell.number_format = fmt
        row += 1


# ========================================================= Sale & Returns


def _returns_sheet(sheet: Worksheet, result: UnderwritingResult, assumptions: UnderwritingAssumptions, report: dict) -> None:
    _widths(sheet, {"A": 40, "B": 18, "C": 4, "D": 10, "E": 16, "F": 16, "G": 16, "H": 16, "I": 16, "J": 18})
    row = _title(sheet, "Sale and returns", "Disposition on forward NOI, then levered and unlevered return attribution.")
    row = _section(sheet, row, "Disposition", "A", "B")
    disposition = result.disposition
    rows = [
        (f"Forward NOI (Year {assumptions.operating.hold_period_years + 1})", result.terminal_noi, MONEY),
        ("Exit cap rate", assumptions.exit_cap_rate, PERCENT),
        ("Gross sale price", disposition.gross_sale_price, MONEY),
        ("Less: selling costs", -disposition.selling_costs, MONEY),
        ("Net sale proceeds before debt", disposition.net_sale_proceeds_before_debt, MONEY),
        ("Less: loan payoff", -disposition.loan_payoff, MONEY),
        ("Net sale proceeds to equity", disposition.net_sale_proceeds_to_equity, MONEY),
        ("Gross sale price vs purchase price", disposition.gross_sale_price / result.purchase_price - 1, PERCENT),
        ("Implied annual appreciation", _annualized(disposition.gross_sale_price / result.purchase_price, assumptions.operating.hold_period_years), PERCENT),
    ]
    for offset, (label, value, fmt) in enumerate(rows):
        _kv(sheet, row + offset, "A", "B", label, value, fmt, bold=label.startswith("Net sale proceeds to equity"))
    attribution_row = row - 1
    _section(sheet, attribution_row, "Return attribution", "D", "J")
    operating_total = sum(result.annual_operating_cash_flows, ZERO)
    sale_gain = disposition.net_sale_proceeds_to_equity - result.initial_equity
    attribution = [
        ("Total equity invested", result.initial_equity, MONEY),
        ("Recurring cash flow over hold", operating_total, MONEY),
        ("Net sale proceeds to equity", disposition.net_sale_proceeds_to_equity, MONEY),
        ("Total distributions", result.returns.total_equity_distributed, MONEY),
        ("Levered profit", result.levered_profit, MONEY),
        ("Profit from operations", operating_total, MONEY),
        ("Profit from sale (proceeds − equity)", sale_gain, MONEY),
        ("Share of profit from sale", _share_of_profit_from_sale(result), PERCENT),
    ]
    for offset, (label, value, fmt) in enumerate(attribution):
        _kv(sheet, attribution_row + 1 + offset, "D", "J", label, value, fmt, bold=label in {"Levered profit", "Total distributions"})
    row = max(row + len(rows), attribution_row + 1 + len(attribution)) + 1

    row = _section(sheet, row, "Levered versus unlevered", "A", "J")
    headers = ["Metric", "Levered", "", "", "Unlevered"]
    for index, header in enumerate(headers, start=1):
        if header:
            _header_cell(sheet.cell(row=row, column=index, value=header))
    row += 1
    for label, levered, unlevered, fmt in (
        ("Capital invested", result.initial_equity, result.unlevered.total_cost, MONEY),
        ("IRR", result.returns.levered_irr, result.unlevered.irr, PERCENT),
        ("Multiple", result.returns.equity_multiple, result.unlevered.multiple, MULTIPLE),
        ("Profit", result.levered_profit, result.unlevered.profit, MONEY),
        ("Leverage effect on IRR", (result.returns.levered_irr - result.unlevered.irr) if result.unlevered.irr is not None else None, None, PERCENT),
    ):
        sheet.cell(row=row, column=1, value=label).font = LABEL_FONT
        cell = sheet.cell(row=row, column=2, value=_num(levered) if levered is not None else "N/A")
        cell.number_format = fmt
        cell = sheet.cell(row=row, column=5, value=_num(unlevered) if unlevered is not None else "")
        cell.number_format = fmt
        row += 1
    row += 1

    row = _section(sheet, row, "Annual return metrics", "A", "J")
    headers = ["Year", "Cash flow after debt", "Cash-on-cash", "DSCR", "Debt yield", "Levered cash flow", "Cumulative levered cash flow", "Unlevered cash flow", "Payback"]
    for index, header in enumerate(headers, start=1):
        _header_cell(sheet.cell(row=row, column=index, value=header))
    row += 1
    cumulative = -result.initial_equity
    paid_back = False
    for index, year in enumerate(result.operating_projection):
        cumulative += result.annual_cash_flows[index]
        payback = ""
        if not paid_back and cumulative >= ZERO:
            payback = "Equity returned"
            paid_back = True
        values = [
            year.year,
            result.annual_operating_cash_flows[index],
            result.returns.cash_on_cash_by_year[index],
            result.returns.dscr_by_year[index] if result.returns.dscr_by_year else None,
            result.annual_debt_yield[index] if result.annual_debt_yield else None,
            result.annual_cash_flows[index],
            cumulative,
            result.unlevered.annual_cash_flows[index],
            payback,
        ]
        formats = [INTEGER, MONEY, PERCENT_1, MULTIPLE, PERCENT_1, MONEY, MONEY, MONEY, None]
        for column, (value, fmt) in enumerate(zip(values, formats), start=1):
            cell = sheet.cell(row=row, column=column, value=_num(value) if isinstance(value, (Decimal, int)) else (value if value is not None else "N/A"))
            if fmt:
                cell.number_format = fmt
        row += 1


def _annualized(multiple: Decimal, years: int) -> Decimal:
    if multiple <= ZERO:
        return ZERO
    return Decimal(str(float(multiple) ** (1 / years) - 1))


# ============================================================ Sensitivity


def _sensitivity_sheet(sheet: Worksheet, result: UnderwritingResult, assumptions: UnderwritingAssumptions, report: dict) -> None:
    _widths(sheet, {"A": 30, **{get_column_letter(i): 14 for i in range(2, 9)}})
    row = _title(sheet, "Sensitivity", f"Every grid holds all other inputs at the analysis case ({_money(result.purchase_price)}). Base case is outlined.")
    exit_caps = _steps(assumptions.exit_cap_rate, Decimal("0.0025"), 2, floor=Decimal("0.005"))
    rent_growth = _steps(assumptions.operating.rent_growth_rate, Decimal("0.005"), 2)
    rates = _steps(assumptions.annual_interest_rate, Decimal("0.005"), 2, floor=ZERO)
    vacancies = _steps(assumptions.operating.vacancy_rate, Decimal("0.01"), 2, floor=ZERO)

    def irr_of(values: dict) -> Decimal | None:
        try:
            return calculate_underwriting(apply_scenario_values(assumptions, values)).returns.levered_irr
        except ValueError:
            return None

    def multiple_of(values: dict) -> Decimal | None:
        try:
            return calculate_underwriting(apply_scenario_values(assumptions, values)).returns.equity_multiple
        except ValueError:
            return None

    def dscr_of(values: dict) -> Decimal | None:
        try:
            return calculate_underwriting(apply_scenario_values(assumptions, values)).returns.minimum_dscr
        except ValueError:
            return None

    target = report.get("target_levered_irr")
    row = _grid(sheet, row, "Levered IRR — exit cap rate (rows) × rent growth (columns)", exit_caps, rent_growth, "exit_cap_rate", "rent_growth_rate", irr_of, PERCENT_1, assumptions, highlight_at=Decimal(str(target)) if target is not None else None)
    row = _grid(sheet, row, "Levered IRR — exit cap rate (rows) × interest rate (columns)", exit_caps, rates, "exit_cap_rate", "annual_interest_rate", irr_of, PERCENT_1, assumptions, highlight_at=Decimal(str(target)) if target is not None else None)
    row = _grid(sheet, row, "Levered IRR — vacancy (rows) × rent growth (columns)", vacancies, rent_growth, "vacancy_rate", "rent_growth_rate", irr_of, PERCENT_1, assumptions, highlight_at=Decimal(str(target)) if target is not None else None)
    row = _grid(sheet, row, "Equity multiple — exit cap rate (rows) × rent growth (columns)", exit_caps, rent_growth, "exit_cap_rate", "rent_growth_rate", multiple_of, MULTIPLE, assumptions)
    if result.loan_amount > ZERO:
        floor = report.get("minimum_dscr")
        row = _grid(sheet, row, "Minimum DSCR — interest rate (rows) × vacancy (columns)", rates, vacancies, "annual_interest_rate", "vacancy_rate", dscr_of, MULTIPLE, assumptions, highlight_at=Decimal(str(floor)) if floor is not None else None)

    if target is not None and report.get("minimum_dscr") is not None:
        target_d = Decimal(str(target))
        floor_d = Decimal(str(report["minimum_dscr"]))
        targets = _steps(target_d, Decimal("0.01"), 2, floor=Decimal("0.0001"))

        def price_of(values: dict) -> Decimal | None:
            case = apply_scenario_values(assumptions, {key: value for key, value in values.items() if key != "target_levered_irr"})
            try:
                return solve_maximum_purchase_price(MaximumPriceAssumptions(case, values["target_levered_irr"], floor_d)).maximum_purchase_price
            except ValueError:
                return None

        row = _grid(sheet, row, "Maximum supportable price — exit cap rate (rows) × target levered IRR (columns)", exit_caps, targets, "exit_cap_rate", "target_levered_irr", price_of, MONEY, assumptions, base_column_value=target_d)

        row = _section(sheet, row, "One-way maximum price sensitivity (re-solved at the same targets)", "A", "E")
        one_way = calculate_one_way_sensitivities(SensitivityAssumptions(assumptions, target_d, floor_d))
        for header_index, header in enumerate(("Driver", "Case", "Input", "Maximum price", "Change vs base"), start=1):
            _header_cell(sheet.cell(row=row, column=header_index, value=header))
        row += 1
        for dimension in one_way.dimensions:
            base = next((point.maximum_purchase_price for point in dimension.points if point.label.startswith("Base")), None)
            for point in dimension.points:
                sheet.cell(row=row, column=1, value=dimension.label)
                sheet.cell(row=row, column=2, value=point.label)
                cell = sheet.cell(row=row, column=3, value=_num(point.value) if point.value is not None else "Scheduled")
                cell.number_format = PERCENT
                cell = sheet.cell(row=row, column=4, value=_num(point.maximum_purchase_price) if point.maximum_purchase_price is not None else (point.error or "Unavailable"))
                cell.number_format = MONEY
                if point.maximum_purchase_price is not None and base is not None:
                    cell = sheet.cell(row=row, column=5, value=_num(point.maximum_purchase_price - base))
                    cell.number_format = MONEY
                row += 1
            row += 1


def _steps(base: Decimal, step: Decimal, count: int, floor: Decimal | None = None) -> list[Decimal]:
    values = [base + step * offset for offset in range(-count, count + 1)]
    if floor is not None:
        values = [max(value, floor) for value in values]
    unique: list[Decimal] = []
    for value in values:
        if value not in unique:
            unique.append(value)
    return unique


def _grid(sheet, row, title, row_values, column_values, row_key, column_key, metric, fmt, assumptions, highlight_at=None, base_column_value=None):
    row = _section(sheet, row, title, "A", get_column_letter(len(column_values) + 1))
    sheet.cell(row=row, column=1, value=f"{_label(row_key)} ↓ / {_label(column_key)} →").font = HEADER_FONT
    for index, value in enumerate(column_values, start=2):
        cell = sheet.cell(row=row, column=index, value=_num(value))
        cell.number_format = PERCENT
        _header_cell(cell)
    row += 1
    base_row_value = _base_of(assumptions, row_key)
    base_col_value = base_column_value if base_column_value is not None else _base_of(assumptions, column_key)
    first_data_row = row
    for row_value in row_values:
        cell = sheet.cell(row=row, column=1, value=_num(row_value))
        cell.number_format = PERCENT
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        for index, column_value in enumerate(column_values, start=2):
            value = metric({row_key: row_value, column_key: column_value})
            cell = sheet.cell(row=row, column=index, value=_num(value) if value is not None else "n/a")
            cell.number_format = fmt
            cell.alignment = Alignment(horizontal="right")
            if row_value == base_row_value and column_value == base_col_value:
                cell.font = TOTAL_FONT
                cell.border = Border(top=Side(style="medium", color=GREEN), bottom=Side(style="medium", color=GREEN), left=Side(style="medium", color=GREEN), right=Side(style="medium", color=GREEN))
        row += 1
    last_column = get_column_letter(len(column_values) + 1)
    data_range = f"B{first_data_row}:{last_column}{row - 1}"
    if highlight_at is not None:
        sheet.conditional_formatting.add(data_range, CellIsRule(operator="lessThan", formula=[str(highlight_at)], fill=PatternFill("solid", fgColor=RED_SOFT)))
        sheet.conditional_formatting.add(data_range, CellIsRule(operator="greaterThanOrEqual", formula=[str(highlight_at)], fill=PatternFill("solid", fgColor=GREEN_SOFT)))
    else:
        sheet.conditional_formatting.add(data_range, ColorScaleRule(start_type="min", start_color="F8D7DA", mid_type="percentile", mid_value=50, mid_color="FFFFFF", end_type="max", end_color="C6E8D0"))
    return row + 1


def _base_of(assumptions: UnderwritingAssumptions, key: str) -> Decimal | None:
    from backend.workflows.underwriting.scenarios import base_value

    try:
        return base_value(assumptions, key)
    except AttributeError:
        return None


def _label(key: str) -> str:
    return {
        "exit_cap_rate": "Exit cap",
        "rent_growth_rate": "Rent growth",
        "annual_interest_rate": "Interest rate",
        "vacancy_rate": "Vacancy",
        "expense_growth_rate": "Expense growth",
        "target_levered_irr": "Target IRR",
    }.get(key, key)


# =================================================================== Risk


def _risk_sheet(sheet: Worksheet, simulation, report: dict) -> None:
    _widths(sheet, {"A": 34, **{get_column_letter(i): 14 for i in range(2, 12)}})
    row = _title(sheet, "Risk simulation", f"{simulation.iterations:,} seeded cases (seed {simulation.seed}); correlated draws through a Gaussian copula with each driver's own distribution.")
    row = _section(sheet, row, "Sampled drivers", "A", "F")
    for index, header in enumerate(("Driver", "Distribution", "Minimum", "Mode", "Maximum", "Confidence"), start=1):
        _header_cell(sheet.cell(row=row, column=index, value=header))
    row += 1
    confidence = {variable.key: f"{variable.confidence_level} / {variable.confidence_percentage}%" for variable in simulation.variables}
    for band in simulation.sampled_ranges:
        sheet.cell(row=row, column=1, value=_label(band.key))
        sheet.cell(row=row, column=2, value=band.distribution)
        for offset, value in enumerate((band.minimum, band.mode, band.maximum), start=3):
            cell = sheet.cell(row=row, column=offset, value=_num(value))
            cell.number_format = PERCENT
        sheet.cell(row=row, column=6, value=confidence.get(band.key, ""))
        row += 1
    row += 1

    row = _section(sheet, row, "Outcome distributions", "A", "K")
    headers = ("Metric", "Minimum", "P5", "P10", "P25", "Median", "P75", "P90", "P95", "Maximum", "Mean")
    for index, header in enumerate(headers, start=1):
        _header_cell(sheet.cell(row=row, column=index, value=header))
    row += 1
    summaries = [
        ("Levered IRR", simulation.levered_irr, PERCENT_1),
        ("Equity multiple", simulation.equity_multiple, MULTIPLE),
        ("Minimum DSCR", simulation.minimum_dscr, MULTIPLE),
        ("Year 1 cash-on-cash", simulation.year_one_cash_on_cash, PERCENT_1),
        ("Terminal NOI", simulation.terminal_noi, MONEY),
        ("Net sale proceeds to equity", simulation.net_sale_proceeds, MONEY),
    ]
    for label, summary, fmt in summaries:
        if summary is None:
            continue
        sheet.cell(row=row, column=1, value=label).font = LABEL_FONT
        values = (summary.minimum, summary.percentile_5, summary.percentile_10, summary.percentile_25, summary.percentile_50, summary.percentile_75, summary.percentile_90, summary.percentile_95, summary.maximum, summary.mean)
        for offset, value in enumerate(values, start=2):
            cell = sheet.cell(row=row, column=offset, value=_num(value))
            cell.number_format = fmt
        row += 1
    row += 1

    row = _section(sheet, row, "Probabilities", "A", "B")
    probabilities = [
        ("Base case levered IRR", simulation.base_case_irr, PERCENT_1),
        ("Probability of meeting the target IRR", simulation.levered_irr.probability_meeting_target, PERCENT_1),
        ("Probability of a loss (IRR below 0%)", simulation.probability_of_loss, PERCENT_1),
        ("Probability minimum DSCR falls below the floor", simulation.probability_dscr_below_floor, PERCENT_1),
        ("Probability of a negative cash-flow year", simulation.probability_negative_cash_flow_year, PERCENT_1),
        ("IRR standard deviation", simulation.levered_irr.standard_deviation, PERCENT_1),
    ]
    for offset, (label, value, fmt) in enumerate(probabilities):
        _kv(sheet, row + offset, "A", "B", label, value if value is not None else "n/a", fmt)
    histogram_row = row - 1
    _section(sheet, histogram_row, "Levered IRR histogram", "D", "G")
    for index, header in enumerate(("From", "To", "Cases", "Share"), start=4):
        _header_cell(sheet.cell(row=histogram_row + 1, column=index, value=header))
    for index, bucket in enumerate(simulation.irr_histogram):
        r = histogram_row + 2 + index
        sheet.cell(row=r, column=4, value=_num(bucket.lower)).number_format = PERCENT_1
        sheet.cell(row=r, column=5, value=_num(bucket.upper)).number_format = PERCENT_1
        sheet.cell(row=r, column=6, value=bucket.count).number_format = INTEGER
        sheet.cell(row=r, column=7, value=f"=F{r}/{simulation.iterations}").number_format = PERCENT_1
    row = max(row + len(probabilities), histogram_row + 2 + len(simulation.irr_histogram)) + 1

    row = _section(sheet, row, "Driver impact (tornado data)", "A", "I")
    headers = ("Driver", "Share of IRR variance", "IRR per +100 bps", "Recurring CF per +100 bps", "Sampled P10", "IRR at P10", "Sampled P90", "IRR at P90", "IRR swing")
    for index, header in enumerate(headers, start=1):
        _header_cell(sheet.cell(row=row, column=index, value=header))
    row += 1
    for impact in simulation.driver_impacts:
        values = (impact.relative_influence, impact.irr_change_per_100_bps, impact.cash_flow_change_per_100_bps, impact.sampled_low, impact.irr_at_low, impact.sampled_high, impact.irr_at_high)
        formats = (PERCENT_1, PERCENT, MONEY, PERCENT, PERCENT_1, PERCENT, PERCENT_1)
        sheet.cell(row=row, column=1, value=_label(impact.key)).font = LABEL_FONT
        for offset, (value, fmt) in enumerate(zip(values, formats), start=2):
            cell = sheet.cell(row=row, column=offset, value=_num(value) if value is not None else "n/a")
            cell.number_format = fmt
        if impact.irr_at_low is not None and impact.irr_at_high is not None:
            cell = sheet.cell(row=row, column=9, value=_num(abs(impact.irr_at_high - impact.irr_at_low)))
            cell.number_format = PERCENT_1
        row += 1
    row += 1

    row = _section(sheet, row, "Representative joint cases", "A", "D")
    for case in (simulation.downside_case, simulation.upside_case):
        sheet.cell(row=row, column=1, value=case.label).font = TOTAL_FONT
        cell = sheet.cell(row=row, column=2, value=_num(case.levered_irr))
        cell.number_format = PERCENT_1
        cell.font = TOTAL_FONT
        row += 1
        for key, value in case.values.items():
            sheet.cell(row=row, column=1, value="   " + _label(key)).font = LABEL_FONT
            cell = sheet.cell(row=row, column=2, value=_num(value))
            cell.number_format = PERCENT
            row += 1
        row += 1
    if simulation.correlations:
        row = _section(sheet, row, "Driver correlations used", "A", "C")
        for correlation in simulation.correlations:
            sheet.cell(row=row, column=1, value=f"{_label(correlation.first_key)} / {_label(correlation.second_key)}")
            sheet.cell(row=row, column=2, value=_num(correlation.coefficient)).number_format = RATIO
            row += 1
        row += 1
    if simulation.market_context:
        context = simulation.market_context
        note = sheet.cell(
            row=row,
            column=1,
            value=f"Calibrated to {context.property_type} · {context.value_bucket} from {context.source} (as of {context.source_as_of_date}; {context.observations} periods, {context.period_start} to {context.period_end}).",
        )
        note.font = NOTE_FONT


# ================================================================ Detail


def _detail_sheet(sheet: Worksheet, result: UnderwritingResult, assumptions: UnderwritingAssumptions, detail: dict) -> None:
    _widths(sheet, {"A": 30, "B": 12, "C": 16, "D": 18, "E": 18, "F": 18, "G": 18, "H": 16, "I": 40})
    row = _title(sheet, "Rent roll and expenses", "Supporting detail saved with this version; totals reconcile to the model inputs.")
    property_detail = detail.get("property") or {}
    if property_detail:
        row = _section(sheet, row, "Property", "A", "C")
        for label, key in (("Name", "name"), ("Address", "address"), ("Class", "property_class"), ("Units", "units"), ("Rentable square feet", "rentable_square_feet"), ("Year built", "year_built")):
            if property_detail.get(key) is not None:
                _kv(sheet, row, "A", "B", label, _num(property_detail[key]) if key in {"units", "rentable_square_feet", "year_built"} else property_detail[key], INTEGER if key in {"units", "rentable_square_feet"} else None)
                row += 1
        row += 1

    rent_roll = detail.get("rent_roll") or []
    if rent_roll:
        row = _section(sheet, row, "Rent roll (in place)", "A", "I")
        headers = ("Unit type", "Units", "Avg SF", "Current rent / mo", "Market rent / mo", "Annual in-place rent", "Annual market rent", "Loss to lease", "Rent / SF / mo")
        for index, header in enumerate(headers, start=1):
            _header_cell(sheet.cell(row=row, column=index, value=header))
        row += 1
        first = row
        for line in rent_roll:
            units = Decimal(str(line["units"]))
            current = Decimal(str(line["current_monthly_rent"]))
            market = Decimal(str(line["market_monthly_rent"])) if line.get("market_monthly_rent") is not None else None
            sf = Decimal(str(line["average_square_feet"])) if line.get("average_square_feet") is not None else None
            sheet.cell(row=row, column=1, value=line["unit_type"])
            sheet.cell(row=row, column=2, value=_num(units)).number_format = INTEGER
            sheet.cell(row=row, column=3, value=_num(sf) if sf else None).number_format = INTEGER
            sheet.cell(row=row, column=4, value=_num(current)).number_format = MONEY
            sheet.cell(row=row, column=5, value=_num(market) if market is not None else None).number_format = MONEY
            sheet.cell(row=row, column=6, value=f"=B{row}*D{row}*12").number_format = MONEY
            sheet.cell(row=row, column=7, value=f"=IF(E{row}=\"\",F{row},B{row}*E{row}*12)").number_format = MONEY
            sheet.cell(row=row, column=8, value=f"=IF(G{row}=0,0,1-F{row}/G{row})").number_format = PERCENT_1
            sheet.cell(row=row, column=9, value=f"=IF(C{row}=\"\",\"\",D{row}/C{row})").number_format = MONEY_CENTS
            row += 1
        sheet.cell(row=row, column=1, value="Total / weighted").font = TOTAL_FONT
        sheet.cell(row=row, column=2, value=f"=SUM(B{first}:B{row - 1})").number_format = INTEGER
        sheet.cell(row=row, column=3, value=f"=IF(B{row}=0,\"\",SUMPRODUCT(B{first}:B{row - 1},C{first}:C{row - 1})/B{row})").number_format = INTEGER
        sheet.cell(row=row, column=4, value=f"=IF(B{row}=0,0,SUMPRODUCT(B{first}:B{row - 1},D{first}:D{row - 1})/B{row})").number_format = MONEY
        sheet.cell(row=row, column=6, value=f"=SUM(F{first}:F{row - 1})").number_format = MONEY
        sheet.cell(row=row, column=7, value=f"=SUM(G{first}:G{row - 1})").number_format = MONEY
        sheet.cell(row=row, column=8, value=f"=IF(G{row}=0,0,1-F{row}/G{row})").number_format = PERCENT_1
        for column in range(1, 10):
            sheet.cell(row=row, column=column).font = TOTAL_FONT
            sheet.cell(row=row, column=column).border = TOP_BORDER
        row += 2

    other_income = detail.get("other_income") or []
    if other_income:
        row = _section(sheet, row, "Other income", "A", "I")
        for index, header in enumerate(("Category", "", "Annual amount", "Per unit", "Note"), start=1):
            if header:
                _header_cell(sheet.cell(row=row, column=index, value=header))
        row += 1
        units = property_detail.get("units")
        first = row
        for line in other_income:
            sheet.cell(row=row, column=1, value=line["category"])
            sheet.cell(row=row, column=3, value=_num(line["annual_amount"])).number_format = MONEY
            if units:
                sheet.cell(row=row, column=4, value=f"=C{row}/{units}").number_format = MONEY
            sheet.cell(row=row, column=5, value=line.get("note") or "")
            row += 1
        sheet.cell(row=row, column=1, value="Total other income").font = TOTAL_FONT
        cell = sheet.cell(row=row, column=3, value=f"=SUM(C{first}:C{row - 1})")
        cell.number_format = MONEY
        cell.font = TOTAL_FONT
        cell.border = TOP_BORDER
        row += 2

    expenses = detail.get("operating_expenses") or []
    if expenses:
        row = _section(sheet, row, "Operating expenses (current year)", "A", "I")
        for index, header in enumerate(("Category", "", "Annual amount", "Per unit", "Per SF", "% of EGI", "% of total", "", "Note"), start=1):
            if header:
                _header_cell(sheet.cell(row=row, column=index, value=header))
        row += 1
        units = property_detail.get("units")
        square_feet = property_detail.get("rentable_square_feet")
        year_one_egi = result.operating_projection[0].effective_gross_income
        first = row
        for line in expenses:
            sheet.cell(row=row, column=1, value=line["category"])
            sheet.cell(row=row, column=3, value=_num(line["annual_amount"])).number_format = MONEY
            if units:
                sheet.cell(row=row, column=4, value=f"=C{row}/{units}").number_format = MONEY
            if square_feet:
                sheet.cell(row=row, column=5, value=f"=C{row}/{square_feet}").number_format = MONEY_CENTS
            sheet.cell(row=row, column=6, value=f"=C{row}/{_num(year_one_egi)}").number_format = PERCENT_1
            sheet.cell(row=row, column=9, value=line.get("note") or "")
            row += 1
        total_row = row
        for r in range(first, total_row):
            sheet.cell(row=r, column=7, value=f"=C{r}/$C${total_row}").number_format = PERCENT_1
        sheet.cell(row=row, column=1, value="Total operating expenses").font = TOTAL_FONT
        for column, formula, fmt in ((3, f"=SUM(C{first}:C{row - 1})", MONEY), (4, f"=SUM(D{first}:D{row - 1})" if units else None, MONEY), (5, f"=SUM(E{first}:E{row - 1})" if square_feet else None, MONEY_CENTS), (6, f"=SUM(F{first}:F{row - 1})", PERCENT_1)):
            if formula:
                cell = sheet.cell(row=row, column=column, value=formula)
                cell.number_format = fmt
                cell.font = TOTAL_FONT
                cell.border = TOP_BORDER
        row += 1
        entered = assumptions.operating.current_operating_expenses
        cell = sheet.cell(row=row, column=1, value="Model input: current operating expenses")
        cell.font = NOTE_FONT
        cell = sheet.cell(row=row, column=3, value=_num(entered))
        cell.number_format = MONEY
        cell.font = NOTE_FONT
        row += 1


# ================================================================= Flags


def _flags_sheet(sheet: Worksheet, result: UnderwritingResult) -> None:
    _widths(sheet, {"A": 28, "B": 12, "C": 70, "D": 14, "E": 14})
    row = _title(sheet, "Risk flags", "Deterministic rules with their thresholds, evaluated at the analysis price.")
    for index, header in enumerate(("Code", "Severity", "Message", "Observed", "Threshold"), start=1):
        _header_cell(sheet.cell(row=row, column=index, value=header))
    row += 1
    if not result.risk_flags:
        sheet.cell(row=row, column=1, value="No rules-based risk flags were triggered.")
        return
    for flag in result.risk_flags:
        sheet.cell(row=row, column=1, value=flag.code)
        sheet.cell(row=row, column=2, value=flag.severity.title())
        sheet.cell(row=row, column=3, value=flag.message)
        fmt = MULTIPLE if "dscr" in flag.code else (MONEY if flag.code == "negative_equity_cash_flow" else PERCENT)
        sheet.cell(row=row, column=4, value=_num(flag.observed_value)).number_format = fmt
        sheet.cell(row=row, column=5, value=_num(flag.threshold)).number_format = fmt
        row += 1


# ============================================================ Assumptions


def _assumptions_sheet(sheet: Worksheet, assumptions: UnderwritingAssumptions, report: dict, detail: dict) -> None:
    _widths(sheet, {"A": 40, "B": 22, "C": 60})
    row = _title(sheet, "Assumptions", "Every input behind the workbook, in model units (rates as decimals).")
    operating = assumptions.operating
    schedule = operating.annual_rent_growth_rates
    groups = (
        ("Version", (
            ("Deal", report.get("deal_name"), None, ""),
            ("Version", report.get("version_label"), None, ""),
            ("Source or document reference", report.get("source"), None, ""),
            ("Effective date", report.get("effective_date"), None, ""),
            ("Confidence", str(report["confidence_level"]).title() if report.get("confidence_level") else None, None, ""),
            ("Analysis basis", report.get("analysis_basis"), None, ""),
        )),
        ("Acquisition and exit", (
            ("Analysis purchase price", assumptions.purchase_price, MONEY, "Price the workbook describes."),
            ("Closing costs", assumptions.closing_cost_rate, PERCENT, "Share of price; funded with equity."),
            ("Hold period", operating.hold_period_years, INTEGER, "Years."),
            ("Exit cap rate", assumptions.exit_cap_rate, PERCENT, "Applied to the year after the hold."),
            ("Selling costs", assumptions.selling_cost_rate, PERCENT, "Share of gross sale price."),
        )),
        ("Operations", (
            ("Current potential gross income", operating.current_potential_gross_income, MONEY, "Annual."),
            ("Current operating expenses", operating.current_operating_expenses, MONEY, "Annual."),
            ("Vacancy and credit loss", operating.vacancy_rate, PERCENT, "Share of potential gross income."),
            ("Rent growth", operating.rent_growth_rate, PERCENT, "Flat annual rate." + (" Overridden by the schedule below." if schedule else "")),
            ("Rent growth schedule", ", ".join(f"{float(rate):.2%}" for rate in schedule) if schedule else "Not supplied", None, "Applies from the first growth year; the last rate continues."),
            ("Expense growth", operating.expense_growth_rate, PERCENT, "Annual."),
            ("Year 1 convention", "Grow once before Year 1 (trailing inputs)" if operating.apply_growth_in_year_one else "Year 1 = current figures; growth starts Year 2", None, ""),
        )),
        ("Financing", (
            ("Loan-to-value", assumptions.loan_to_value, PERCENT, "0 for all-cash."),
            ("Interest rate", assumptions.annual_interest_rate, PERCENT, "Annual note rate, monthly compounding."),
            ("Amortization", assumptions.amortization_years, INTEGER, "Years."),
        )),
        ("Targets", (
            ("Target levered IRR", report.get("target_levered_irr"), PERCENT, ""),
            ("Minimum DSCR", report.get("minimum_dscr"), MULTIPLE, ""),
            ("Maximum supportable price", report.get("maximum_purchase_price"), MONEY, "Solved from the targets."),
            ("Binding constraint", (report.get("binding_constraint") or "").replace("_", " ") or None, None, ""),
        )),
    )
    for title, rows in groups:
        row = _section(sheet, row, title, "A", "C")
        for label, value, fmt, note in rows:
            if value is None:
                continue
            _kv(sheet, row, "A", "B", label, value, fmt)
            sheet.cell(row=row, column=3, value=note).font = NOTE_FONT
            row += 1
        row += 1
    if detail.get("notes"):
        row = _section(sheet, row, "Analyst notes", "A", "C")
        cell = sheet.cell(row=row, column=1, value=detail["notes"])
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        sheet.row_dimensions[row].height = 60


# ================================================================ Helpers


def _title(sheet: Worksheet, title: str, subtitle: str) -> int:
    sheet.cell(row=1, column=1, value=title).font = TITLE_FONT
    sheet.cell(row=2, column=1, value=subtitle).font = SUBTITLE_FONT
    return 4


def _section(sheet: Worksheet, row: int, title: str, start: str, end: str) -> int:
    cell = sheet[f"{start}{row}"]
    cell.value = title
    cell.font = SECTION_FONT
    cell.fill = SECTION_FILL
    if start != end:
        sheet.merge_cells(f"{start}{row}:{end}{row}")
    return row + 1


def _header_cell(cell) -> None:
    cell.font = HEADER_FONT
    cell.fill = HEADER_FILL
    cell.border = BOTTOM_BORDER
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _kv(sheet: Worksheet, row: int, label_column: str, value_column: str, label, value, fmt, *, bold=False) -> None:
    label_cell = sheet[f"{label_column}{row}"]
    label_cell.value = label
    label_cell.font = TOTAL_FONT if bold else LABEL_FONT
    if value is None:
        return
    cell = sheet[f"{value_column}{row}"]
    cell.value = _num(value) if isinstance(value, (Decimal, int, float)) else value
    if fmt and isinstance(value, (Decimal, int, float)):
        cell.number_format = fmt
    cell.alignment = Alignment(horizontal="right")
    cell.font = TOTAL_FONT if bold else LABEL_FONT


def _widths(sheet: Worksheet, widths: dict[str, int]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width


def _negatives_red(sheet: Worksheet, cell_range: str) -> None:
    sheet.conditional_formatting.add(cell_range, CellIsRule(operator="lessThan", formula=["0"], font=Font(color="B91C1C")))


def _num(value):
    """Excel stores doubles; convert Decimal without carrying string precision."""
    if isinstance(value, Decimal):
        return float(value)
    return value


def _money(value) -> str:
    return f"${float(value):,.0f}"


def _convention_notes(assumptions: UnderwritingAssumptions) -> list[str]:
    convention = (
        "Year 1 equals the entered current income and expenses; growth begins in Year 2."
        if not assumptions.operating.apply_growth_in_year_one
        else "Entered figures are treated as trailing and grown once before Year 1."
    )
    return [
        convention + " Going-in cap rate is Year 1 NOI divided by the analysis price.",
        "The sale is capitalized on forward (year after hold) NOI less selling costs; the loan balance is repaid from proceeds. Closing costs are funded with equity, not financed.",
        "Debt amortizes monthly at the stated note rate; annual debt service is the sum of twelve payments. Cash-on-cash excludes sale proceeds; IRR and multiple include them.",
        "Pro Forma subtotals and return metrics are live Excel formulas over engine values; the engine's own figures are shown beside them for reference. Educational modeling output, not financial advice.",
    ]
